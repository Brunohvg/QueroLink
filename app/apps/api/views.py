from datetime import date, timedelta
import logging

from rest_framework import viewsets, status, generics, serializers as drf_serializers
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from django.core.exceptions import ObjectDoesNotExist, ValidationError
from django.utils import timezone
from django.db import transaction
from django.db.models import Max, Sum, Q
from django.http import HttpResponse
from django.conf import settings
from django_ratelimit.decorators import ratelimit
from django.utils.decorators import method_decorator
from rest_framework_simplejwt.views import TokenObtainPairView
from drf_spectacular.utils import extend_schema, OpenApiParameter, OpenApiTypes

from app.apps.sales.models import Sale, SaleChangeLog
from app.apps.sellers.models import Seller
from app.apps.commissions.models import (
    CommissionPeriod,
    SellerCommission,
    CommissionAdjustment,
)
from app.apps.commissions.services import (
    calculate_estimated_commission,
    get_commission_rate,
    get_period_by_legacy_label,
)
from app.apps.accounts.models import User

from .serializers import (
    SellerSerializer,
    SellerCreateSerializer,
    SellerImportSerializer,
    SaleSerializer,
    SaleCreateSerializer,
    CommissionPeriodSerializer,
    CommissionPeriodCreateSerializer,
    ChangePasswordSerializer,
    ManagerSaleUpdateSerializer,
    SaleChangeLogSerializer,
    SellerDayJustificationSerializer,
    SellerDayJustificationCreateSerializer,
    SellerDayJustificationUpdateSerializer,
)
from .permissions import IsManagerOrAdmin, IsFinancialOrAdmin, IsSellerOwner
from app.apps.audit.utils import log_action

logger = logging.getLogger(__name__)


class SellerViewSet(viewsets.ModelViewSet):
    queryset = Seller.objects.select_related('user').all()
    permission_classes = [IsAuthenticated, IsManagerOrAdmin]

    def get_serializer_class(self):
        if self.action == 'create':
            return SellerCreateSerializer
        return SellerSerializer

    def get_queryset(self):
        return Seller.objects.filter(
            tenant=self.request.user.tenant,
        ).select_related('user')

    @action(detail=False, methods=['get'], url_path='dashboard-summary')
    def dashboard_summary(self, request):
        """Dados da tabela de vendedores em lote, sem fan-out por vendedor."""
        from decimal import Decimal, ROUND_HALF_UP

        from app.apps.commissions.day_status import get_period_summary_bulk
        from app.apps.commissions.services import (
            PeriodIntegrityError,
            PeriodNotFound,
            get_commission_rate,
            resolve_selected_period,
        )

        tenant = request.user.tenant
        sellers = list(
            Seller.objects.filter(tenant=tenant)
            .select_related('tenant', 'user')
            .order_by('name')
        )

        try:
            period = resolve_selected_period(request, tenant)
        except PeriodNotFound:
            return Response({'error': 'Competencia nao encontrada.'}, status=404)
        except PeriodIntegrityError as exc:
            return Response({'error': str(exc)}, status=409)

        sales_by_seller = {}
        commissions_by_seller = {}
        day_summaries = {}
        if period:
            sales_by_seller = {
                row['seller_id']: row
                for row in Sale.objects.filter(
                    tenant=tenant,
                    seller_id__in=[seller.pk for seller in sellers],
                    origin__in=Sale.COMMISSION_ORIGINS,
                    status='ATIVA',
                    sale_date__gte=period.start_date,
                    sale_date__lte=period.end_date,
                ).values('seller_id').annotate(
                    month_total=Sum('amount'),
                    last_sale_date=Max('sale_date'),
                )
            }
            commissions_by_seller = {
                commission.seller_id: commission
                for commission in SellerCommission.objects.filter(
                    period=period,
                    seller_id__in=[seller.pk for seller in sellers],
                )
            }
            day_summaries = get_period_summary_bulk(
                tenant, period, sellers, reference_date=timezone.localdate(),
            )

        today = timezone.localdate()
        result = []
        for seller in sellers:
            sales = sales_by_seller.get(seller.pk, {})
            month_total = sales.get('month_total') or 0
            last_sale_date = sales.get('last_sale_date')
            commission = commissions_by_seller.get(seller.pk)
            summary = day_summaries.get(seller.pk, {})

            if (
                commission
                and period.status != CommissionPeriod.Status.ABERTA
            ):
                commission_amount = commission.commission_amount
            else:
                rate = get_commission_rate(seller)
                commission_amount = int(
                    (Decimal(month_total) * rate).quantize(
                        Decimal('1'), rounding=ROUND_HALF_UP,
                    )
                )

            launched = summary.get('lancados', 0)
            justified = summary.get('justificados', 0)
            pending = summary.get('pendentes', 0)
            if launched == 0 and justified == 0:
                operational_status = 'SEM_LANCAMENTO'
            else:
                operational_status = summary.get(
                    'operational_status', 'PENDENTE',
                )

            result.append({
                **SellerSerializer(seller).data,
                'month_total': month_total,
                'commission_amount': commission_amount,
                'submitted_days': launched + justified,
                'expected_days': launched + justified + pending,
                'operational_status': operational_status,
                'financial_status': (
                    commission.status if commission else 'ABERTA'
                ),
                'last_sale_date': (
                    last_sale_date.isoformat() if last_sale_date else None
                ),
                'has_sale_today': last_sale_date == today,
            })

        return Response({
            'period_uuid': str(period.uuid) if period else None,
            'sellers': result,
        })

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        result = serializer.save()
        seller = self.get_queryset().get(uuid=result['uuid'])
        log_action(request, 'seller.created', instance=seller)

        from app.apps.accounts.models import mark_onboarding_step
        mark_onboarding_step(request.user.tenant, 'step_sellers')

        return Response(result, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=['post'])
    def reset_password(self, request, pk=None):
        seller = self.get_object()
        if not seller.user:
            return Response(
                {'error': 'Vendedor sem usuario vinculado.'}, status=400,
            )
        from app.apps.accounts.utils import generate_temp_password
        password = generate_temp_password()
        seller.user.set_password(password)
        seller.user.save()
        try:
            from app.apps.notifications.tasks import notify_seller_credentials
            notify_seller_credentials(seller, password)
            whatsapp_sent = True
            whatsapp_error = None
        except Exception as e:
            whatsapp_sent = False
            whatsapp_error = str(e)
            logger.error(
                'WhatsApp failed on reset_password for seller %s: %s',
                seller.uuid, e, exc_info=True
            )
        log_action(request, 'seller.password_reset', instance=seller)
        return Response({
            'message': 'Senha redefinida com sucesso.',
            'whatsapp_sent': whatsapp_sent,
        })

    @action(detail=False, methods=['post'])
    def import_sellers(self, request):
        serializer = SellerImportSerializer(data=request.data, context={'request': request})
        serializer.is_valid(raise_exception=True)
        result = serializer.save()
        log_action(request, 'seller.import', changes={
            'total': result['total'],
            'criados': result['criados'],
            'erros': result['erros'],
        })
        return Response(result, status=status.HTTP_200_OK)

    def perform_destroy(self, instance):
        if Sale.objects.filter(seller=instance).exists():
            from rest_framework import serializers as drf_ser
            raise drf_ser.ValidationError({
                'detail': 'Este vendedor possui vendas registradas e nao pode ser excluido. '
                          'Desative-o para preservar o historico.',
            })
        log_action(self.request, 'seller.deleted', instance=instance,
                   changes={'name': instance.name, 'username': instance.user.username if instance.user else None})
        user = instance.user
        instance.delete()
        if user:
            user.delete()


class SaleViewSet(viewsets.ModelViewSet):
    permission_classes = [IsAuthenticated, IsSellerOwner | IsManagerOrAdmin]

    def get_serializer_class(self):
        if self.action in ('create', 'update', 'partial_update'):
            return SaleCreateSerializer
        return SaleSerializer

    def get_queryset(self):
        user = self.request.user
        qs = Sale.objects.select_related('seller', 'order').filter(
            tenant=user.tenant,
        )
        if user.role == User.Role.SELLER:
            qs = qs.filter(seller=user.seller_profile)
        return qs

    def perform_create(self, serializer):
        sale = serializer.save()
        log_action(self.request, 'sale.created', instance=sale)

        from app.apps.commissions.services import ensure_seller_commission
        ensure_seller_commission(sale.seller, sale.sale_date)

        from app.apps.accounts.models import mark_onboarding_step
        mark_onboarding_step(self.request.user.tenant, 'step_first_sale')

    def perform_update(self, serializer):
        serializer.save(updated_by=self.request.user)
        log_action(self.request, 'sale.updated', instance=serializer.instance)

    def perform_destroy(self, instance):
        from app.apps.commissions.services import validate_sale_can_be_changed
        can_change, error_msg = validate_sale_can_be_changed(
            instance.seller, instance.sale_date, self.request.user,
        )
        if not can_change:
            raise drf_serializers.ValidationError({'detail': error_msg})
        log_action(self.request, 'sale.deleted', instance=instance)
        instance.delete()

    @action(detail=False, methods=['post'], permission_classes=[IsAuthenticated, IsManagerOrAdmin], url_path='import/preview')
    def import_preview(self, request):
        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'error': 'Arquivo obrigatorio.'}, status=400)

        if file_obj.size > 5 * 1024 * 1024:
            return Response({'error': 'Arquivo muito grande. Maximo 5MB.'}, status=400)

        name = file_obj.name.lower()
        content_bytes = file_obj.read()

        from app.apps.sales.services import (
            _parse_csv, _parse_xlsx, normalize_headers, preview_import_rows,
        )
        from hashlib import sha256
        from app.apps.sales.models import SaleImportBatch

        file_hash = sha256(content_bytes).hexdigest()

        if name.endswith('.csv'):
            content = None
            for _enc in ('utf-8-sig', 'cp1252', 'latin-1'):
                try:
                    content = content_bytes.decode(_enc)
                    break
                except UnicodeDecodeError:
                    continue
            if content is None:
                content = content_bytes.decode('utf-8', errors='replace')
            rows = _parse_csv(content)
        elif name.endswith('.xlsx'):
            import io
            rows = _parse_xlsx(content_bytes)
            if rows is None:
                return Response({'error': 'Arquivo XLSX invalido.'}, status=400)
        else:
            return Response({'error': 'Formato nao suportado. Envie CSV ou XLSX.'}, status=400)

        if not rows or len(rows) < 2:
            return Response({'error': 'Arquivo vazio ou sem dados.'}, status=400)

        headers = rows[0]
        data_rows = rows[1:]
        header_map = normalize_headers(headers)

        if 'date' not in header_map or 'seller' not in header_map or 'amount' not in header_map:
            missing = []
            if 'date' not in header_map:
                missing.append('data')
            if 'seller' not in header_map:
                missing.append('vendedor')
            if 'amount' not in header_map:
                missing.append('valor')
            return Response({
                'error': f'Cabecalhos obrigatorios nao encontrados: {", ".join(missing)}.',
            }, status=400)

        if len(data_rows) > 500:
            return Response({'error': 'Maximo 500 linhas por importacao.'}, status=400)

        tenant = request.user.tenant
        results = preview_import_rows(tenant, data_rows, header_map)

        dupe_batch = SaleImportBatch.objects.filter(
            tenant=tenant, file_hash=file_hash,
        ).first()
        already_imported = dupe_batch is not None

        return Response({
            'filename': file_obj.name,
            'file_hash': file_hash,
            'total_rows': len(data_rows),
            'headers': headers,
            'results': results,
            'already_imported': already_imported,
            'ok_count': sum(1 for r in results if r['status'] == 'ok'),
            'duplicate_count': sum(1 for r in results if r['status'] == 'duplicate'),
            'error_count': sum(1 for r in results if r['status'] == 'error'),
            'needs_selection_count': sum(1 for r in results if r['status'] == 'needs_selection'),
        })

    @action(detail=False, methods=['post'], permission_classes=[IsAuthenticated, IsManagerOrAdmin], url_path='import/confirm')
    def import_confirm(self, request):
        confirmed = request.data.get('rows', [])
        filename = request.data.get('filename', '')
        file_hash = request.data.get('file_hash', '')
        force = request.data.get('force_reimport', False)

        if not confirmed:
            return Response({'error': 'Nenhuma linha para importar.'}, status=400)
        if not filename or not file_hash:
            return Response({'error': 'filename e file_hash obrigatorios.'}, status=400)

        tenant = request.user.tenant

        from app.apps.sales.models import SaleImportBatch
        existing = SaleImportBatch.objects.filter(
            tenant=tenant, file_hash=file_hash,
        ).first()
        if existing and not force:
            return Response({
                'error': 'Este arquivo ja foi importado. Use force_reimport para confirmar.',
                'existing_batch_uuid': str(existing.uuid),
            }, status=409)

        from app.apps.sales.services import import_sales as do_import

        try:
            result = do_import(
                tenant, request.user, confirmed, filename, file_hash,
            )
        except ValueError as e:
            return Response({'error': str(e)}, status=400)

        log_action(
            request, 'sales_import',
            changes={
                'filename': filename,
                'file_hash': file_hash,
                'created': result['created'],
                'duplicates': result['duplicates'],
                'errors_count': len(result['errors']),
            },
        )

        return Response(result, status=200)

    @action(detail=True, methods=['patch'], permission_classes=[IsAuthenticated, IsManagerOrAdmin], url_path='manager-update')
    def manager_update(self, request, pk=None):
        sale = self.get_object()
        serializer = ManagerSaleUpdateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        data = serializer.validated_data
        reason = data.pop('reason')

        from app.apps.sales.services import update_sale_as_manager

        try:
            sale, changed = update_sale_as_manager(sale, request.user, data, reason)
        except ValueError as e:
            return Response({'error': str(e)}, status=status.HTTP_409_CONFLICT)

        if not changed:
            return Response({'message': 'Nenhuma alteração detectada.', 'changed': False})

        return Response({
            'message': 'Venda atualizada com sucesso.',
            'changed': True,
            'sale': SaleSerializer(sale).data,
        })

    @action(detail=True, methods=['get'], permission_classes=[IsAuthenticated, IsSellerOwner | IsManagerOrAdmin], url_path='history')
    def history(self, request, pk=None):
        sale = self.get_object()
        qs = SaleChangeLog.objects.filter(
            sale=sale, tenant=request.user.tenant,
        ).select_related('changed_by').order_by('-changed_at')
        serializer = SaleChangeLogSerializer(qs, many=True)
        return Response(serializer.data)


class CommissionPeriodViewSet(viewsets.ModelViewSet):
    permission_classes = [IsAuthenticated, IsManagerOrAdmin]

    EDITABLE_STATUSES = (
        CommissionPeriod.Status.ABERTA,
        CommissionPeriod.Status.PARCIALMENTE_FECHADA,
        CommissionPeriod.Status.PARCIALMENTE_PAGA,
    )

    def get_serializer_class(self):
        if self.action == 'create':
            return CommissionPeriodCreateSerializer
        return CommissionPeriodSerializer

    def get_queryset(self):
        qs = CommissionPeriod.objects.filter(
            tenant=self.request.user.tenant,
        )

        scope = self.request.query_params.get('scope')
        if scope == 'active':
            qs = qs.exclude(status__in=[
                CommissionPeriod.Status.PAGA,
                CommissionPeriod.Status.CANCELADA,
            ])
        elif scope == 'finished':
            qs = qs.filter(status__in=[
                CommissionPeriod.Status.PAGA,
                CommissionPeriod.Status.CANCELADA,
            ])

        month = self.request.query_params.get('month')
        if month:
            qs = qs.filter(month=int(month))
        year = self.request.query_params.get('year')
        if year:
            qs = qs.filter(year=int(year))

        return qs.prefetch_related('seller_commissions__seller')

    def list(self, request, *args, **kwargs):
        from app.apps.commissions.services import sync_period_seller_commissions
        editable = self.get_queryset().filter(status__in=self.EDITABLE_STATUSES)[:6]
        for period in editable:
            try:
                sync_period_seller_commissions(period)
            except Exception:
                logger.exception('Auto-sync falhou para period %s', period.pk)
        return super().list(request, *args, **kwargs)

    def retrieve(self, request, *args, **kwargs):
        period = self.get_object()
        if period.status in self.EDITABLE_STATUSES:
            from app.apps.commissions.services import sync_period_seller_commissions
            sync_period_seller_commissions(period)
        return super().retrieve(request, *args, **kwargs)

    def perform_create(self, serializer):
        tenant = self.request.user.tenant
        period = serializer.save(tenant=tenant)

        from app.apps.commissions.services import sync_period_seller_commissions
        sync_period_seller_commissions(period)

    @action(detail=True, methods=['post'])
    def sync(self, request, pk=None):
        from app.apps.commissions.services import sync_period_seller_commissions
        period = self.get_object()
        created = sync_period_seller_commissions(period)
        return Response({
            'synced': True,
            'created': created,
            'message': 'Competencia sincronizada com sucesso.' if created > 0
                       else 'Competencia ja estava sincronizada.',
        })

    @action(detail=True, methods=['post'])
    @method_decorator(ratelimit(key='user', rate='30/h', method='POST', block=True))
    def close_sellers(self, request, pk=None):
        from app.apps.commissions.services import close_seller_commissions

        period = self.get_object()
        seller_ids = request.data.get('seller_commission_ids', [])

        if not seller_ids:
            return Response(
                {'error': 'Selecione ao menos um vendedor.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            calculations = close_seller_commissions(period, seller_ids, request.user)
        except ValueError as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)

        log_action(
            request, 'commission_period.close_sellers', instance=period,
            changes={
                'month': period.month, 'year': period.year,
                'seller_ids': seller_ids,
                'closed_count': len(calculations),
            },
        )
        return Response({
            'closed': len(calculations),
            'commissions': calculations,
        })

    @action(
        detail=True, methods=['post'],
        permission_classes=[IsAuthenticated, IsManagerOrAdmin],
    )
    @method_decorator(ratelimit(key='user', rate='30/h', method='POST', block=True))
    def reopen_sellers(self, request, pk=None):
        from app.apps.commissions.services import reopen_seller_commissions

        period = self.get_object()
        seller_ids = request.data.get('seller_commission_ids', [])
        reason = request.data.get('reason', '').strip()

        if not seller_ids:
            return Response(
                {'error': 'Selecione ao menos um vendedor.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            commissions = reopen_seller_commissions(period, seller_ids, request.user, reason)
        except ValueError as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)

        log_action(
            request, 'commission_period.reopen_sellers', instance=period,
            changes={
                'month': period.month, 'year': period.year,
                'seller_ids': seller_ids,
                'reason': reason,
            },
        )
        return Response({
            'reopened': len(commissions),
            'message': f'{len(commissions)} vendedor(es) reaberto(s).',
        })

    @action(
        detail=True, methods=['post'],
        permission_classes=[IsAuthenticated, IsManagerOrAdmin],
    )
    @method_decorator(ratelimit(key='user', rate='10/h', method='POST', block=True))
    def send_accounting(self, request, pk=None):
        period = self.get_object()
        tenant = request.user.tenant

        if period.tenant_id != tenant.pk:
            return Response(
                {'error': 'Competencia nao pertence ao seu tenant.'},
                status=status.HTTP_403_FORBIDDEN,
            )

        if not tenant.accountant_email:
            return Response(
                {'error': 'Cadastre o e-mail do contador em Configuracoes.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if period.status == CommissionPeriod.Status.CANCELADA:
            return Response(
                {'error': 'Competencia cancelada nao pode ser enviada.'},
                status=status.HTTP_409_CONFLICT,
            )

        force_resend = request.data.get('force_resend', False)
        if period.sent_to_accounting_at and not force_resend:
            return Response(
                {'error': 'Competencia ja foi enviada. Use force_resend para reenviar.'},
                status=status.HTTP_409_CONFLICT,
            )

        scs = list(period.seller_commissions.all())
        if not scs:
            return Response(
                {'error': 'Nenhum vendedor sincronizado nesta competencia.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        open_statuses = [
            SellerCommission.Status.ABERTA,
            SellerCommission.Status.REABERTA,
        ]
        has_open = any(sc.status in open_statuses for sc in scs)
        if has_open:
            return Response(
                {'error': 'Existem vendedores com comissao em aberto. Feche todas antes de enviar.'},
                status=status.HTTP_409_CONFLICT,
            )

        from app.apps.notifications.tasks import send_accounting_package_email

        send_accounting_package_email.delay(
            str(tenant.uuid),
            period.month,
            period.year,
            requested_by_user_id=str(request.user.pk),
        )

        log_action(
            request, 'commission_period.send_accounting', instance=period,
            changes={
                'month': period.month, 'year': period.year,
                'force_resend': force_resend,
            },
        )

        return Response({
            'message': 'Envio agendado com sucesso.',
            'detail': f'Pacote contabil sera enviado para {tenant.accountant_email}.',
        }, status=status.HTTP_202_ACCEPTED)

    @action(
        detail=True, methods=['post'],
        permission_classes=[IsAuthenticated, IsFinancialOrAdmin],
    )
    @method_decorator(ratelimit(key='user', rate='30/h', method='POST', block=True))
    def pay_sellers(self, request, pk=None):
        from app.apps.commissions.services import pay_seller_commissions

        period = self.get_object()
        seller_ids = request.data.get('seller_commission_ids', [])

        if not seller_ids:
            return Response(
                {'error': 'Selecione ao menos um vendedor.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        payment_data = {
            'payment_date': request.data.get('payment_date'),
            'payment_method': request.data.get('payment_method', ''),
            'payment_notes': request.data.get('payment_notes', ''),
        }

        try:
            commissions = pay_seller_commissions(period, seller_ids, request.user, payment_data)
        except ValueError as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)

        log_action(
            request, 'commission_period.pay_sellers', instance=period,
            changes={
                'month': period.month, 'year': period.year,
                'seller_ids': seller_ids,
                'payment_date': payment_data.get('payment_date'),
            },
        )
        return Response({
            'paid': len(commissions),
            'message': f'{len(commissions)} vendedor(es) pago(s).',
        })

    def update(self, request, *args, **kwargs):
        from app.apps.commissions.services import update_period

        partial = kwargs.pop('partial', False)
        period = self.get_object()
        data = request.data

        try:
            period, changed = update_period(period, data, request.user)
        except ValueError as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)

        log_action(
            request, 'commission_period.updated', instance=period,
            changes={
                'month': period.month, 'year': period.year,
                'changed_fields': changed,
                'new_data': dict(data),
            },
        )
        return Response(self.get_serializer(period).data)

    def perform_destroy(self, instance):
        from app.apps.commissions.services import delete_period

        try:
            sc_count = delete_period(instance, self.request.user)
        except ValueError as e:
            from rest_framework import serializers as drf_ser
            raise drf_ser.ValidationError({'detail': str(e)})

        log_action(
            self.request, 'commission_period.deleted', instance=instance,
            changes={
                'month': instance.month, 'year': instance.year,
                'seller_commissions_removed': sc_count,
            },
        )

    @action(
        detail=True, methods=['post'],
        permission_classes=[IsAuthenticated, IsManagerOrAdmin],
    )
    def cancel(self, request, pk=None):
        from app.apps.commissions.services import cancel_period

        period = self.get_object()
        reason = request.data.get('reason', '').strip()

        try:
            period = cancel_period(period, reason, request.user)
        except ValueError as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)

        log_action(
            request, 'commission_period.cancelled', instance=period,
            changes={
                'month': period.month, 'year': period.year,
                'reason': reason,
            },
        )
        return Response(self.get_serializer(period).data)

    @action(
        detail=True, methods=['get'],
        permission_classes=[IsAuthenticated, IsFinancialOrAdmin | IsManagerOrAdmin],
        url_path='receipt/(?P<sc_id>[^/.]+)',
    )
    def receipt(self, request, pk=None, sc_id=None):
        from django.template.loader import render_to_string
        from weasyprint import HTML

        period = self.get_object()
        try:
            sc = SellerCommission.objects.select_related(
                'seller', 'period',
            ).prefetch_related('adjustments').get(
                pk=sc_id, period=period,
            )
        except SellerCommission.DoesNotExist:
            return Response({'error': 'Comissao nao encontrada.'}, status=404)

        if sc.status != SellerCommission.Status.PAGA:
            return Response({'error': 'Recibo disponivel apenas para comissoes pagas.'}, status=404)

        adjustments = sc.adjustments.all()
        hoje = timezone.now()
        html = render_to_string('reports/recibo_comissao.html', {
            'tenant': period.tenant,
            'period': period,
            'sc': sc,
            'adjustments': adjustments,
            'hoje': hoje,
        })
        pdf = HTML(string=html).write_pdf()
        return HttpResponse(pdf, content_type='application/pdf')

    @action(
        detail=True, methods=['get'],
        permission_classes=[IsAuthenticated, IsFinancialOrAdmin | IsManagerOrAdmin],
        url_path='receipts',
    )
    def receipts(self, request, pk=None):
        from django.template.loader import render_to_string
        from weasyprint import HTML

        period = self.get_object()
        commissions = SellerCommission.objects.filter(
            period=period,
            status=SellerCommission.Status.PAGA,
        ).select_related('seller').prefetch_related('adjustments').order_by('seller__name')

        if not commissions.exists():
            return Response({'error': 'Nenhuma comissao paga no periodo.'}, status=404)

        hoje = timezone.now()
        pages = []
        for sc in commissions:
            adjustments = sc.adjustments.all()
            html = render_to_string('reports/recibo_comissao.html', {
                'tenant': period.tenant,
                'period': period,
                'sc': sc,
                'adjustments': adjustments,
                'hoje': hoje,
            })
            pages.append(html)

        full_html = ''.join(
            f'<div style="page-break-after: always;">{p}</div>' if i < len(pages) - 1 else p
            for i, p in enumerate(pages)
        )
        pdf = HTML(string=full_html).write_pdf()
        return HttpResponse(pdf, content_type='application/pdf')

    @action(detail=False, methods=['get'], permission_classes=[IsAuthenticated, IsManagerOrAdmin], url_path='suggest')
    def suggest(self, request):
        from app.apps.commissions.services import (
            suggest_period_range, suggest_next_open_month,
        )

        tenant = request.user.tenant
        month_param = request.query_params.get('month')
        year_param = request.query_params.get('year')

        if month_param is None and year_param is None:
            month, year = suggest_next_open_month(tenant)
        else:
            hoje = timezone.localdate()
            try:
                month = int(month_param) if month_param is not None else hoje.month
                year = int(year_param) if year_param is not None else hoje.year
            except (ValueError, TypeError):
                return Response({'error': 'Mes/ano invalidos.'}, status=status.HTTP_400_BAD_REQUEST)

        if month < 1 or month > 12:
            return Response({'error': 'Mes deve estar entre 1 e 12.'}, status=status.HTTP_400_BAD_REQUEST)
        if year < 2000 or year > 2100:
            return Response({'error': 'Ano invalido.'}, status=status.HTTP_400_BAD_REQUEST)

        start, end = suggest_period_range(tenant, month, year)
        return Response({
            'month': month,
            'year': year,
            'period_start_day': tenant.period_start_day,
            'start_date': start.isoformat(),
            'end_date': end.isoformat(),
        })

    @action(detail=False, methods=['get'], permission_classes=[IsAuthenticated, IsManagerOrAdmin], url_path='preview')
    def preview(self, request):
        from datetime import date
        from app.apps.commissions.services import (
            calculate_estimated_commission, get_missing_days_before_today,
        )

        tenant = request.user.tenant
        hoje = timezone.localdate()
        month = int(request.query_params.get('month', hoje.month))
        year = int(request.query_params.get('year', hoje.year))

        sellers = Seller.objects.filter(tenant=tenant, is_active=True)
        sellers_data = []
        total_sold = 0
        total_commission = 0

        for seller in sellers:
            est_commission, est_total = calculate_estimated_commission(seller, month, year)
            try:
                missing_days = get_missing_days_before_today(seller, month, year)
            except Exception:
                missing_days = []
            sellers_data.append({
                'name': seller.name,
                'uuid': str(seller.uuid),
                'total_sold': est_total,
                'commission_rate': float(get_commission_rate(seller)),
                'commission_amount': est_commission,
                'has_missing_days': len(missing_days) > 0,
                'missing_days_count': len(missing_days),
            })
            total_sold += est_total
            total_commission += est_commission

        sellers_data.sort(key=lambda s: s['total_sold'], reverse=True)

        return Response({
            'sellers': sellers_data,
            'totals': {
                'total_sold': total_sold,
                'total_commission': total_commission,
            },
            'month': month,
            'year': year,
        })

@extend_schema(
    responses={200: dict},
    description='Retorna períodos com comissões fechadas/ajustadas para pagamento.',
)
class PaymentQueueView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated, IsFinancialOrAdmin | IsManagerOrAdmin]

    def get(self, request):
        tenant = request.user.tenant
        periods = CommissionPeriod.objects.filter(
            tenant=tenant,
        ).exclude(
            status=CommissionPeriod.Status.CANCELADA,
        ).prefetch_related(
            'seller_commissions__seller',
        ).order_by('-year', '-month')

        result = []
        for period in periods:
            fechadas = [
                sc for sc in period.seller_commissions.all()
                if sc.status in (SellerCommission.Status.FECHADA, SellerCommission.Status.AJUSTADA)
            ]
            if not fechadas:
                continue

            from .serializers import SellerCommissionReadSerializer
            result.append({
                'period': {
                    'uuid': str(period.uuid),
                    'month': period.month,
                    'year': period.year,
                    'status': period.status,
                },
                'commissions': SellerCommissionReadSerializer(fechadas, many=True).data,
            })

        return Response(result)


class JWTLoginView(TokenObtainPairView):
    @method_decorator(ratelimit(key='ip', rate='5/m', method='POST', block=True))
    @method_decorator(
        ratelimit(key='post:username', rate='5/m', method='POST', block=True),
    )
    def post(self, request, *args, **kwargs):
        return super().post(request, *args, **kwargs)


class SellerSalesListView(generics.ListAPIView):
    serializer_class = SaleSerializer
    permission_classes = [IsAuthenticated, IsSellerOwner]

    def get_queryset(self):
        try:
            seller = self.request.user.seller_profile
        except ObjectDoesNotExist:
            return Sale.objects.none()
        return Sale.objects.filter(
            seller=seller, tenant=self.request.user.tenant,
        )


@extend_schema(
    parameters=[
        OpenApiParameter('month', int, required=False),
        OpenApiParameter('year', int, required=False),
    ],
    responses={200: dict},
)
class RankingView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated, IsManagerOrAdmin]

    def get(self, request):
        from app.apps.commissions.services import (
            calculate_estimated_commission,
            calculate_estimated_commission_for_period,
            get_commission_rate,
            get_period_by_legacy_label,
            legacy_month_range,
        )
        from decimal import Decimal, ROUND_HALF_UP

        tenant = request.user.tenant
        month = int(request.query_params.get(
            'month', timezone.localdate().month,
        ))
        year = int(request.query_params.get(
            'year', timezone.localdate().year,
        ))
        if month < 1 or month > 12:
            return Response({'error': 'Mes invalido (1-12).'}, status=400)
        if year < 2020:
            return Response({'error': 'Ano invalido.'}, status=400)
        period = get_period_by_legacy_label(tenant, month, year)
        if period:
            start = period.start_date
            end = period.end_date
        else:
            start, end = legacy_month_range(month, year)

        sales = Sale.objects.filter(
            tenant=tenant,
            origin__in=Sale.COMMISSION_ORIGINS,
            status='ATIVA',
            sale_date__gte=start,
            sale_date__lte=end,
        ).values('seller__uuid', 'seller__name', 'seller__commission_rate').annotate(
            total_sold=Sum('amount'),
            sale_count=Sum(1),
        ).order_by('-total_sold')

        ranking = []
        seller_uuids = [s['seller__uuid'] for s in sales]
        sellers_map = {
            str(s.uuid): s
            for s in Seller.objects.filter(uuid__in=seller_uuids, tenant=tenant)
        }

        for s in sales:
            seller_uuid = s['seller__uuid']
            total_sold = s['total_sold']
            sale_count = s['sale_count']
            ticket_medio = round(total_sold / sale_count) if sale_count > 0 else 0

            seller_obj = sellers_map.get(str(seller_uuid))
            if seller_obj:
                rate = get_commission_rate(seller_obj)
                if period:
                    commission_estimada, _ = calculate_estimated_commission_for_period(seller_obj, period)
                else:
                    commission_estimada = int((Decimal(str(total_sold)) * Decimal(str(rate))).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
            else:
                rate = Decimal('0')
                commission_estimada = 0

            ranking.append({
                'seller__uuid': seller_uuid,
                'seller__name': s['seller__name'],
                'total_sold': total_sold,
                'sale_count': sale_count,
                'ticket_medio': ticket_medio,
                'commission_rate': float(rate),
                'commission_estimada': commission_estimada,
            })

        top_seller = ranking[0] if ranking else None

        return Response({
            'month': month,
            'year': year,
            'ranking': ranking,
            'top_seller': top_seller,
        })


class CommissionPeriodsByStatusView(generics.ListAPIView):
    serializer_class = CommissionPeriodSerializer
    permission_classes = [IsAuthenticated, IsManagerOrAdmin | IsFinancialOrAdmin]

    def get_queryset(self):
        tenant = self.request.user.tenant
        status_filter = self.kwargs.get('status', 'ABERTA')
        return CommissionPeriod.objects.filter(
            tenant=tenant, status=status_filter,
        ).prefetch_related(
            'seller_commissions__seller',
        ).order_by('-year', '-month')


@extend_schema(responses={(200, 'text/csv'): OpenApiTypes.BINARY})
class CommissionPeriodCsvView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated, IsFinancialOrAdmin | IsManagerOrAdmin]

    def get(self, request, pk=None):
        import csv
        import io

        period = CommissionPeriod.objects.get(
            uuid=pk, tenant=request.user.tenant,
        )
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([
            'Vendedor', 'Total Vendido (R$)', 'Taxa (%)',
            'Comissao (R$)', 'Status',
        ])
        for sc in period.seller_commissions.select_related('seller').all():
            writer.writerow([
                sc.seller.name,
                f'{sc.total_sold_amount / 100:.2f}',
                f'{float(sc.commission_rate) * 100:.2f}',
                f'{sc.commission_amount / 100:.2f}',
                sc.get_status_display(),
            ])
        response = HttpResponse(
            buf.getvalue(), content_type='text/csv; charset=utf-8',
        )
        response['Content-Disposition'] = (
            f'attachment; filename="comissao_{period.month}_{period.year}.csv"'
        )
        return response


class ManagerSalesListView(generics.ListAPIView):
    serializer_class = SaleSerializer
    permission_classes = [IsAuthenticated, IsManagerOrAdmin]

    def get_queryset(self):
        tenant = self.request.user.tenant
        qs = Sale.objects.select_related('seller').filter(tenant=tenant)
        seller_uuid = self.request.query_params.get('seller')
        if seller_uuid:
            qs = qs.filter(seller__uuid=seller_uuid)
        return qs


@extend_schema(
    parameters=[
        OpenApiParameter('seller_id', str, location=OpenApiParameter.PATH),
        OpenApiParameter('start', str, required=False, description='Data inicial (YYYY-MM-DD)'),
        OpenApiParameter('end', str, required=False, description='Data final (YYYY-MM-DD)'),
        OpenApiParameter('month', int, required=False),
        OpenApiParameter('year', int, required=False),
    ],
    responses={200: dict},
)
class SellerDetailView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated, IsManagerOrAdmin]

    def get(self, request, seller_id=None):
        from app.apps.commissions.services import (
            resolve_selected_period,
            calculate_estimated_commission_for_period,
            PeriodNotFound, PeriodIntegrityError,
        )

        tenant = request.user.tenant
        try:
            seller = Seller.objects.select_related('user').get(
                uuid=seller_id, tenant=tenant,
            )
        except Seller.DoesNotExist:
            return Response(
                {'error': 'Vendedor nao encontrado.'}, status=404,
            )

        start_str = request.query_params.get('start')
        end_str = request.query_params.get('end')
        period_param = request.query_params.get('period')
        selected_period = None
        is_custom = False

        if period_param:
            try:
                selected_period = resolve_selected_period(request, tenant)
            except PeriodNotFound:
                return Response({'error': 'Competencia nao encontrada.'}, status=404)
            except PeriodIntegrityError as exc:
                return Response({'error': str(exc)}, status=409)
            start = selected_period.start_date
            end = selected_period.end_date
        elif start_str and end_str:
            is_custom = True
            start = date.fromisoformat(start_str)
            end = date.fromisoformat(end_str)
        else:
            try:
                selected_period = resolve_selected_period(request, tenant)
            except PeriodIntegrityError as exc:
                return Response({'error': str(exc)}, status=409)
            if selected_period:
                start = selected_period.start_date
                end = selected_period.end_date
            else:
                start = end = None

        if start is None:
            manual_sales_qs = Sale.objects.none()
            link_sales_qs = Sale.objects.none()
        else:
            manual_sales_qs = Sale.objects.filter(
                tenant=tenant, seller=seller,
                origin__in=Sale.COMMISSION_ORIGINS,
                sale_date__gte=start, sale_date__lte=end,
            ).order_by('-sale_date', '-created_at')

            link_sales_qs = Sale.objects.filter(
                tenant=tenant, seller=seller,
                origin=Sale.Origin.LINK,
                sale_date__gte=start, sale_date__lte=end,
            ).order_by('-sale_date', '-created_at')

        manual_total = manual_sales_qs.filter(status='ATIVA').aggregate(
            t=Sum('amount'),
        )['t'] or 0
        link_total = link_sales_qs.filter(status='ATIVA').aggregate(
            t=Sum('amount'),
        )['t'] or 0

        manual_sales_qs_slice = list(manual_sales_qs[:200])
        manual_sales = []
        from django.db.models import Count
        log_counts = dict(
            SaleChangeLog.objects.filter(
                sale__in=[s.pk for s in manual_sales_qs_slice],
            ).values('sale_id').annotate(
                count=Count('uuid'),
            ).values_list('sale_id', 'count')
        )
        for s in manual_sales_qs_slice:
            manual_sales.append({
                'uuid': str(s.uuid),
                'amount': s.amount,
                'origin': s.origin,
                'origin_display': s.get_origin_display(),
                'sale_date': s.sale_date.isoformat(),
                'notes': s.notes or '',
                'status': s.status,
                'change_log_count': log_counts.get(s.uuid, 0),
            })

        link_sales = []
        for s in link_sales_qs[:200]:
            link_sales.append({
                'uuid': str(s.uuid),
                'amount': s.amount,
                'origin': s.origin,
                'origin_display': s.get_origin_display(),
                'sale_date': s.sale_date.isoformat(),
                'notes': s.notes or '',
            })

        commissions = SellerCommission.objects.filter(
            seller=seller,
        ).select_related('period').prefetch_related(
            'adjustments',
        ).order_by('-period__year', '-period__month')

        commissions_data = []
        for sc in commissions:
            adjustments_data = [
                {
                    'previous_amount': a.previous_amount,
                    'new_amount': a.new_amount,
                    'difference': a.difference,
                    'reason': a.reason,
                    'created_at': a.created_at.isoformat(),
                }
                for a in sc.adjustments.all()
            ]

            period_status = sc.period.status
            total_sold = sc.total_sold_amount
            commission_amount = sc.commission_amount

            if period_status == CommissionPeriod.Status.ABERTA:
                est, total = calculate_estimated_commission_for_period(
                    seller, sc.period,
                )
                total_sold = total
                commission_amount = est

            commissions_data.append({
                'period_month': sc.period.month,
                'period_year': sc.period.year,
                'period_status': period_status,
                'total_sold_amount': total_sold,
                'commission_rate': float(sc.commission_rate),
                'commission_amount': commission_amount,
                'status': sc.status,
                'operational_status': sc.operational_status,
                'submitted_days_count': sc.submitted_days_count,
                'expected_working_days': sc.expected_working_days,
                'missing_days_count': sc.missing_days_count,
                'closed_at': sc.closed_at.isoformat() if sc.closed_at else None,
                'paid_at': sc.paid_at.isoformat() if sc.paid_at else None,
                'paid_amount': sc.paid_amount,
                'payment_date': sc.payment_date.isoformat() if sc.payment_date else None,
                'payment_method': sc.payment_method,
                'adjustments': adjustments_data,
            })

        evo = manual_sales_qs.filter(status='ATIVA').values(
            'sale_date',
        ).annotate(
            day_total=Sum('amount'),
        ).order_by('sale_date')
        evolution = [
            {'date': e['sale_date'].isoformat(), 'total': e['day_total']}
            for e in evo
        ]

        comp_data = []
        last_periods = list(
            CommissionPeriod.objects.filter(tenant=tenant)
            .exclude(status=CommissionPeriod.Status.CANCELADA)
            .order_by('-start_date')[:6]
        )
        last_periods.reverse()
        for p in last_periods:
            mt = Sale.objects.filter(
                tenant=tenant, seller=seller,
                origin__in=Sale.COMMISSION_ORIGINS, status='ATIVA',
                sale_date__gte=p.start_date, sale_date__lte=p.end_date,
            ).aggregate(t=Sum('amount'))['t'] or 0
            mc = Sale.objects.filter(
                tenant=tenant, seller=seller,
                origin__in=Sale.COMMISSION_ORIGINS, status='ATIVA',
                sale_date__gte=p.start_date, sale_date__lte=p.end_date,
            ).count()
            sc_c = commissions.filter(period=p).first()

            if sc_c and sc_c.period.status != CommissionPeriod.Status.ABERTA:
                commission_value = sc_c.commission_amount
            else:
                commission_value, _ = calculate_estimated_commission_for_period(seller, p)

            comp_data.append({
                'month': p.display_label,
                'total': mt,
                'sale_count': mc,
                'commission': commission_value,
            })

        selected_period_data = None
        selected_period_commission = None
        if selected_period:
            sp_sc = commissions.filter(period=selected_period).first()
            rate = get_commission_rate(seller)
            if sp_sc and sp_sc.period.status != CommissionPeriod.Status.ABERTA:
                sp_total = sp_sc.total_sold_amount
                sp_commission = sp_sc.commission_amount
                sp_rate = float(sp_sc.commission_rate)
                sp_status = sp_sc.status
                sp_is_estimated = False
            else:
                sp_commission, sp_total = calculate_estimated_commission_for_period(
                    seller, selected_period,
                )
                sp_rate = float(rate)
                sp_status = sp_sc.status if sp_sc else 'ABERTA'
                sp_is_estimated = True
            selected_period_data = {
                'uuid': str(selected_period.uuid),
                'label': selected_period.label,
                'display_label': selected_period.display_label,
                'start_date': selected_period.start_date.isoformat(),
                'end_date': selected_period.end_date.isoformat(),
                'status': selected_period.status,
            }
            selected_period_commission = {
                'total_sold_amount': sp_total,
                'commission_amount': sp_commission,
                'commission_rate': sp_rate,
                'status': sp_status,
                'is_estimated': sp_is_estimated,
            }

        # LOTE 5 - controle operacional de dias (justificativas), sem valor
        # financeiro. Apenas quando ha uma competencia selecionada (nao custom).
        day_status = None
        if selected_period is not None:
            from app.apps.commissions.day_status import get_period_day_statuses
            ds = get_period_day_statuses(tenant, seller, selected_period)
            day_status = {
                'summary': ds['summary'],
                'days': [
                    {
                        'date': d['date'].isoformat(),
                        'is_expected_day': d['is_expected_day'],
                        'status': d['status'],
                        'active_sales_count': d['active_sales_count'],
                        'active_sales_total': d['active_sales_total'],
                        'justification_uuid': d['justification_uuid'],
                        'justification_reason': d['justification_reason'],
                        'justification_reason_display': d['justification_reason_display'],
                        'justification_notes': d['justification_notes'],
                        'can_manage_justification': d['can_manage_justification'],
                    }
                    for d in ds['days']
                ],
            }

        return Response({
            'seller': {
                'uuid': str(seller.uuid),
                'name': seller.name,
                'phone': (seller.phone or '')[:4] + '****' + (seller.phone or '')[-4:] if len(seller.phone or '') >= 8 else '****',
                'is_active': seller.is_active,
                'commission_rate': float(get_commission_rate(seller)),
                'cpf': seller.cpf,
                'cpf_formatted': seller.cpf_formatted,
                'created_at': (
                    seller.created_at.isoformat()
                    if seller.created_at else None
                ),
                'username': seller.user.username if seller.user else None,
            },
            'period': (
                {'start': start.isoformat(), 'end': end.isoformat()}
                if start else None
            ),
            'is_custom_range': is_custom,
            'selected_period': selected_period_data,
            'selected_period_commission': selected_period_commission,
            'manual_total': manual_total,
            'link_total': link_total,
            'manual_sale_count': manual_sales_qs.count(),
            'link_sale_count': link_sales_qs.count(),
            'manual_sales': manual_sales,
            'link_sales': link_sales,
            'commissions': commissions_data,
            'evolution': evolution,
            'comparison': comp_data,
            'day_status': day_status,
        })


@extend_schema(responses={(200, 'text/csv'): OpenApiTypes.BINARY})
class SellerReportCsvView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated, IsManagerOrAdmin]

    def get(self, request, seller_id=None):
        import csv
        import io
        from app.apps.commissions.services import resolve_selected_period, PeriodNotFound, PeriodIntegrityError

        tenant = request.user.tenant
        seller = Seller.objects.get(uuid=seller_id, tenant=tenant)

        start_str = request.query_params.get('start')
        end_str = request.query_params.get('end')
        period_param = request.query_params.get('period')
        report_period = None
        is_custom = False

        if period_param:
            try:
                report_period = resolve_selected_period(request, tenant)
            except PeriodNotFound:
                return Response({'error': 'Competencia nao encontrada.'}, status=404)
            except PeriodIntegrityError as exc:
                return Response({'error': str(exc)}, status=409)
            start = report_period.start_date
            end = report_period.end_date
        elif start_str and end_str:
            is_custom = True
            start = date.fromisoformat(start_str)
            end = date.fromisoformat(end_str)
        else:
            try:
                report_period = resolve_selected_period(request, tenant)
            except PeriodIntegrityError as exc:
                return Response({'error': str(exc)}, status=409)
            if not report_period:
                return Response({'detail': 'Nenhuma competencia disponivel.'})
            start = report_period.start_date
            end = report_period.end_date

        sales = Sale.objects.filter(
            tenant=tenant, seller=seller,
            origin__in=Sale.COMMISSION_ORIGINS,
            sale_date__gte=start, sale_date__lte=end,
        ).order_by('sale_date')

        commission = SellerCommission.objects.filter(
            seller=seller, period=report_period,
        ).first() if report_period else None

        commission_amount = 0
        commission_rate_display = 0
        status_display = 'Aberta'
        if commission:
            if commission.period.status == CommissionPeriod.Status.ABERTA:
                from app.apps.commissions.services import calculate_estimated_commission_for_period
                est, _ = calculate_estimated_commission_for_period(seller, commission.period)
                commission_amount = est
            else:
                commission_amount = commission.commission_amount
            commission_rate_display = float(commission.commission_rate)
            status_display = commission.period.get_status_display()

        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(['Data', 'Valor (R$)', 'Observacao'])
        total = 0
        for s in sales:
            writer.writerow([
                s.sale_date.isoformat(),
                f'{s.amount / 100:.2f}',
                s.notes or '',
            ])
            total += s.amount

        writer.writerow([])
        writer.writerow(['TOTAL', f'{total / 100:.2f}', ''])
        if commission:
            writer.writerow([
                'Taxa de comissao',
                f'{commission_rate_display * 100:.2f}%', '',
            ])
            writer.writerow([
                'Comissao calculada',
                f'{commission_amount / 100:.2f}', '',
            ])
            writer.writerow([
                'Status', status_display, '',
            ])
        elif is_custom:
            writer.writerow([
                'Comissão oficial disponível apenas por competência', '', '',
            ])
        writer.writerow([])
        writer.writerow([f'Vendedor: {seller.name}'])
        writer.writerow([f'Empresa: {tenant.company_name}'])
        writer.writerow([f'Periodo: {start.isoformat()} a {end.isoformat()}'])

        response = HttpResponse(
            buf.getvalue(), content_type='text/csv; charset=utf-8',
        )
        response['Content-Disposition'] = (
            f'attachment; filename="{seller.name}_{start}_{end}.csv"'
        )
        return response


@extend_schema(responses={(200, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'): OpenApiTypes.BINARY})
class SellerReportExcelView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated, IsManagerOrAdmin]

    def get(self, request, seller_id=None):
        import io
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
        from app.apps.commissions.services import resolve_selected_period, PeriodNotFound, PeriodIntegrityError

        tenant = request.user.tenant
        seller = Seller.objects.get(uuid=seller_id, tenant=tenant)

        start_str = request.query_params.get('start')
        end_str = request.query_params.get('end')
        period_param = request.query_params.get('period')
        report_period = None
        is_custom = False

        if period_param:
            try:
                report_period = resolve_selected_period(request, tenant)
            except PeriodNotFound:
                return Response({'error': 'Competencia nao encontrada.'}, status=404)
            except PeriodIntegrityError as exc:
                return Response({'error': str(exc)}, status=409)
            start = report_period.start_date
            end = report_period.end_date
        elif start_str and end_str:
            is_custom = True
            start = date.fromisoformat(start_str)
            end = date.fromisoformat(end_str)
        else:
            try:
                report_period = resolve_selected_period(request, tenant)
            except PeriodIntegrityError as exc:
                return Response({'error': str(exc)}, status=409)
            if not report_period:
                return Response({'detail': 'Nenhuma competencia disponivel.'})
            start = report_period.start_date
            end = report_period.end_date

        sales = Sale.objects.filter(
            tenant=tenant, seller=seller,
            origin__in=Sale.COMMISSION_ORIGINS,
            sale_date__gte=start, sale_date__lte=end,
        ).order_by('sale_date')

        commission = SellerCommission.objects.filter(
            seller=seller, period=report_period,
        ).first() if report_period else None

        commission_amount = 0
        commission_rate_display = 0
        status_display = 'Aberta'
        if commission:
            if commission.period.status == CommissionPeriod.Status.ABERTA:
                from app.apps.commissions.services import calculate_estimated_commission_for_period
                est, _ = calculate_estimated_commission_for_period(seller, commission.period)
                commission_amount = est
            else:
                commission_amount = commission.commission_amount
            commission_rate_display = float(commission.commission_rate)
            status_display = commission.period.get_status_display()

        wb = Workbook()
        ws = wb.active
        ws.title = 'Vendas (Comissao)'

        header_font = Font(bold=True, size=12)
        total_font = Font(bold=True, size=11)
        header_fill = PatternFill(
            start_color='4361EE', end_color='4361EE', fill_type='solid',
        )
        header_font_white = Font(bold=True, color='FFFFFF', size=11)

        ws.merge_cells('A1:C1')
        ws['A1'] = f'{seller.name} — {tenant.company_name}'
        ws['A1'].font = header_font
        ws.merge_cells('A2:C2')
        ws['A2'] = (
            f'Periodo: {start.strftime("%d/%m/%Y")} '
            f'a {end.strftime("%d/%m/%Y")}'
        )
        ws['A2'].font = Font(size=10, color='666666')

        ws.append([])
        headers = ['Data', 'Valor (R$)', 'Observacao']
        ws.append(headers)
        for col in range(1, 4):
            cell = ws.cell(row=4, column=col)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

        total = 0
        for s in sales:
            ws.append([
                s.sale_date.strftime('%d/%m/%Y'),
                s.amount / 100,
                s.notes or '',
            ])
            total += s.amount

        ws.append([])
        ws.append(['TOTAL', total / 100, ''])
        for col in range(1, 4):
            ws.cell(row=ws.max_row, column=col).font = total_font

        if commission:
            ws.append([])
            ws.append([
                'Taxa de comissao',
                f'{commission_rate_display * 100:.2f}%', '',
            ])
            ws.append([
                'Comissao calculada',
                commission_amount / 100, '',
            ])
            ws.append([
                'Status',
                status_display, '',
            ])
        elif is_custom:
            ws.append([])
            ws.append([
                'Comissão oficial disponível apenas por competência', '', '',
            ])

        ws.column_dimensions['A'].width = 14
        ws.column_dimensions['B'].width = 18
        ws.column_dimensions['C'].width = 40

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        response = HttpResponse(
            buf.getvalue(),
            content_type=(
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            ),
        )
        response['Content-Disposition'] = (
            f'attachment; filename="{seller.name}_{start}_{end}.xlsx"'
        )
        return response


@extend_schema(responses={(200, 'application/pdf'): OpenApiTypes.BINARY})
class SellerReportPdfView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated, IsManagerOrAdmin]

    def get(self, request, seller_id=None):
        from django.template.loader import render_to_string
        from app.apps.commissions.services import resolve_selected_period, PeriodNotFound, PeriodIntegrityError

        tenant = request.user.tenant
        seller = Seller.objects.get(uuid=seller_id, tenant=tenant)

        start_str = request.query_params.get('start')
        end_str = request.query_params.get('end')
        period_param = request.query_params.get('period')
        hoje = timezone.localdate()
        report_period = None
        is_custom = False

        if period_param:
            try:
                report_period = resolve_selected_period(request, tenant)
            except PeriodNotFound:
                return Response({'error': 'Competencia nao encontrada.'}, status=404)
            except PeriodIntegrityError as exc:
                return Response({'error': str(exc)}, status=409)
            start = report_period.start_date
            end = report_period.end_date
        elif start_str and end_str:
            is_custom = True
            start = date.fromisoformat(start_str)
            end = date.fromisoformat(end_str)
        else:
            try:
                report_period = resolve_selected_period(request, tenant)
            except PeriodIntegrityError as exc:
                return Response({'error': str(exc)}, status=409)
            if not report_period:
                return Response({'detail': 'Nenhuma competencia disponivel.'})
            start = report_period.start_date
            end = report_period.end_date

        sales = Sale.objects.filter(
            tenant=tenant, seller=seller,
            origin__in=Sale.COMMISSION_ORIGINS,
            sale_date__gte=start, sale_date__lte=end,
        ).order_by('sale_date')

        total = sum(s.amount for s in sales)

        commission = SellerCommission.objects.filter(
            seller=seller, period=report_period,
        ).first() if report_period else None

        commission_amount = 0
        commission_rate_display = 0
        status_display = 'Aberta'
        if commission:
            if commission.period.status == CommissionPeriod.Status.ABERTA:
                from app.apps.commissions.services import calculate_estimated_commission_for_period
                est, _ = calculate_estimated_commission_for_period(seller, commission.period)
                commission_amount = est
            else:
                commission_amount = commission.commission_amount
            commission_rate_display = float(commission.commission_rate)
            status_display = commission.period.get_status_display()

        def _fmt(cents):
            r = cents // 100
            c = cents % 100
            return f'{r:,}.{c:02d}'.replace(',', '.')

        def _pct(rate):
            return f'{float(rate) * 100:.2f}%'

        total_fmt = _fmt(total)
        commission_amount_fmt = _fmt(commission_amount)
        rate_fmt = _pct(commission_rate_display)

        html = render_to_string('reports/seller_report_pdf.html', {
            'seller': seller,
            'tenant': tenant,
            'start': start,
            'end': end,
            'sales': sales,
            'total': total,
            'total_fmt': total_fmt,
            'commission': commission,
            'commission_notice': (
                'Comissão oficial disponível apenas por competência'
                if is_custom else ''
            ),
            'report_period': report_period,
            'commission_amount': commission_amount,
            'commission_amount_fmt': commission_amount_fmt,
            'rate_fmt': rate_fmt,
            'status_display': status_display,
            'hoje': hoje,
        })

        from weasyprint import HTML
        pdf = HTML(string=html).write_pdf()

        response = HttpResponse(pdf, content_type='application/pdf')
        response['Content-Disposition'] = (
            f'attachment; filename="{seller.name}_{start}_{end}.pdf"'
        )
        return response


@extend_schema(
    parameters=[OpenApiParameter('year', int, required=False)],
    responses={200: dict},
)
class AnnualRankingView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated, IsManagerOrAdmin]

    def get(self, request):
        tenant = request.user.tenant
        year = int(request.query_params.get(
            'year', timezone.localdate().year,
        ))

        rankings = SellerCommission.objects.filter(
            period__tenant=tenant,
            period__year=year,
        ).values('seller__uuid', 'seller__name').annotate(
            total=Sum('total_sold_amount'),
            commission_total=Sum('commission_amount'),
        ).order_by('-total')[:5]

        return Response({
            'year': year,
            'ranking': list(rankings),
        })


@extend_schema(
    parameters=[
        OpenApiParameter('month', int, required=False),
        OpenApiParameter('year', int, required=False),
    ],
    responses={200: dict},
)
class DashboardSummaryView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated, IsManagerOrAdmin]

    def get(self, request):
        tenant = request.user.tenant
        month = request.query_params.get('month')
        year = request.query_params.get('year')

        from app.apps.commissions.services import (
            get_dashboard_data, resolve_selected_period,
            PeriodNotFound, PeriodIntegrityError,
        )

        period_param = request.query_params.get('period')
        if period_param:
            try:
                period = resolve_selected_period(request, tenant)
            except PeriodNotFound:
                return Response({'error': 'Competencia nao encontrada.'}, status=404)
            except PeriodIntegrityError as exc:
                return Response({'error': str(exc)}, status=409)
            data = get_dashboard_data(tenant, period=period)
            return Response(data)

        if month:
            month = int(month)
        if year:
            year = int(year)

        data = get_dashboard_data(tenant, month=month, year=year)

        return Response(data)


@extend_schema(
    request=dict,
    responses={200: dict, 201: dict, 400: dict, 500: dict},
    methods=['GET', 'POST'],
)
class SellerLinkCreateView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated, IsSellerOwner]

    def get(self, request):
        try:
            seller = request.user.seller_profile
        except ObjectDoesNotExist:
            return Response({'error': 'Perfil de vendedor nao encontrado.'}, status=400)

        from app.apps.orders.models import Order
        orders = Order.objects.filter(
            seller=seller, tenant=seller.tenant,
        ).select_related('seller').prefetch_related(
            'payments',
        ).order_by('-created_at')[:50]

        orders_data = []
        for o in orders:
            try:
                link = o.payment_link
                link_url = link.gateway_url if link else None
            except Exception:
                link_url = None
            payment = o.payments.first()
            refusal = payment.refusal_reason if payment else None
            orders_data.append({
                'uuid': str(o.uuid),
                'customer_name': o.customer_name,
                'total_amount': o.total_amount,
                'status': o.status,
                'status_display': o.get_status_display(),
                'link_url': link_url,
                'refusal_reason': refusal,
                'created_at': o.created_at.isoformat(),
            })

        return Response(orders_data)

    @method_decorator(ratelimit(key='user', rate='10/m', method='POST', block=True))
    def post(self, request):
        try:
            seller = request.user.seller_profile
        except ObjectDoesNotExist:
            return Response(
                {'error': 'Perfil de vendedor nao encontrado.'}, status=400,
            )

        tenant = request.user.tenant
        if not tenant:
            return Response(
                {'error': 'Usuario sem tenant.'}, status=400,
            )

        customer_name = request.data.get('customer_name', '').strip()
        amount_str = request.data.get('amount', '').strip()
        installments = int(request.data.get('installments', 1))

        if not customer_name:
            return Response(
                {'error': 'Nome do cliente e obrigatorio.'}, status=400,
            )
        if not amount_str:
            return Response(
                {'error': 'Valor e obrigatorio.'}, status=400,
            )

        try:
            amount_str = (
                amount_str.replace('R$', '').replace(',', '.').strip()
            )
            amount_cents = int(float(amount_str) * 100)
            if amount_cents <= 0:
                raise ValueError
        except (ValueError, TypeError):
            return Response({'error': 'Valor invalido.'}, status=400)

        if installments < 1 or installments > 12:
            return Response({'error': 'Numero de parcelas invalido (1-12).'}, status=400)

        try:
            if not tenant.pagarme_api_key:
                return Response(
                    {'error': 'Configure a chave do Pagar.me nas configuracoes.'}, status=400,
                )

            from app.apps.orders.services import (
                create_payment_link as make_link,
            )
            order, link_url = make_link(
                tenant=tenant,
                seller=seller,
                customer_name=customer_name,
                amount_cents=amount_cents,
                installments=installments,
            )

            log_action(request, 'order.link_created', instance=order)

            try:
                from app.apps.notifications.tasks import notify_seller_link_status
                notify_seller_link_status(seller, order, 'link_created')
            except Exception:
                logger.error(
                    'Failed to send link_created notification for order %s',
                    order.uuid, exc_info=True
                )

            return Response({
                'uuid': str(order.uuid),
                'customer_name': order.customer_name,
                'total_amount': order.total_amount,
                'link_url': link_url,
            }, status=201)

        except Exception as e:
            import logging
            logging.getLogger(__name__).error("SellerLinkCreateView error: %s", e, exc_info=True)
            return Response({'error': str(e)}, status=500)


class ChangePasswordView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = ChangePasswordSerializer

    def post(self, request):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        serializer.save()
        log_action(request, 'user.password_changed')
        return Response({'message': 'Senha alterada com sucesso.'})


@extend_schema(responses={200: dict})
class WebhookStatusView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated, IsManagerOrAdmin]

    def get(self, request):
        from django.conf import settings
        from app.apps.webhooks.models import WebhookEvent

        tenant = request.user.tenant
        webhook_url = (
            f"https://{settings.SERVICE_FQDN_WEB}"
            f"/api/webhooks/pagarme/{tenant.slug}/"
        )

        last_event = (
            WebhookEvent.objects.filter(
                gateway='pagarme', tenant=tenant,
            )
            .order_by('-received_at')
            .first()
        )

        status_data = {
            'webhook_url': webhook_url,
            'tenant_slug': tenant.slug,
            'last_event': None,
        }

        if last_event:
            status_data['last_event'] = {
                'received_at': last_event.received_at.isoformat(),
                'processed': last_event.processed,
                'error': last_event.processing_error,
            }

        return Response(status_data)


class PushSubscribeView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        try:
            endpoint = request.data['endpoint']
            keys = request.data['keys']
            p256dh = keys['p256dh']
            auth = keys['auth']
        except (KeyError, TypeError):
            return Response({'error': 'Dados invalidos. Envie endpoint, keys.p256dh e keys.auth.'}, status=400)

        from app.apps.notifications.models import PushSubscription

        sub, created = PushSubscription.objects.update_or_create(
            user=request.user,
            endpoint=endpoint,
            defaults={
                'tenant': request.user.tenant,
                'p256dh': p256dh,
                'auth': auth,
                'is_active': True,
            },
        )
        return Response({'status': 'subscribed', 'created': created}, status=201 if created else 200)


class PushUnsubscribeView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        endpoint = request.data.get('endpoint', '')
        if not endpoint:
            return Response({'error': 'endpoint e obrigatorio.'}, status=400)

        from app.apps.notifications.models import PushSubscription

        PushSubscription.objects.filter(
            user=request.user, endpoint=endpoint,
        ).update(is_active=False)
        return Response(status=204)


@extend_schema(responses={(200, 'application/pdf'): OpenApiTypes.BINARY})
class SellerStatementView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, year=None, month=None):
        if request.user.role != request.user.Role.SELLER:
            return Response({'error': 'Apenas vendedores.'}, status=403)

        try:
            seller = request.user.seller_profile
        except ObjectDoesNotExist:
            return Response({'error': 'Perfil de vendedor nao encontrado.'}, status=404)

        from app.apps.commissions.services import (
            resolve_selected_period, get_period_by_legacy_label,
            calculate_estimated_commission_for_period,
            PeriodNotFound, PeriodIntegrityError,
        )

        today = timezone.localdate()
        period_param = request.query_params.get('period')
        statement_period = None
        is_month_fallback = False

        if period_param or (year is None and month is None):
            try:
                statement_period = resolve_selected_period(request, seller.tenant)
            except PeriodNotFound:
                return Response({'error': 'Competencia nao encontrada.'}, status=404)
            except PeriodIntegrityError as exc:
                return Response({'error': str(exc)}, status=409)
            if not statement_period:
                return Response({'detail': 'Nenhuma competencia disponivel.'})
            start = statement_period.start_date
            end = statement_period.end_date
            year_int = statement_period.year
            month_int = statement_period.month
        else:
            try:
                year_int = int(year) if year else today.year
                month_int = int(month) if month else today.month
            except (ValueError, TypeError):
                return Response({'error': 'Mes/ano invalidos.'}, status=400)
            statement_period = get_period_by_legacy_label(seller.tenant, month_int, year_int)
            if statement_period:
                start = statement_period.start_date
                end = statement_period.end_date
            else:
                is_month_fallback = True
                import calendar as _cal
                start = date(year_int, month_int, 1)
                end = date(year_int, month_int, _cal.monthrange(year_int, month_int)[1])

        sales = Sale.objects.filter(
            seller=seller,
            sale_date__gte=start,
            sale_date__lte=end,
        ).order_by('sale_date')

        sales_ativas = [s for s in sales if s.status == 'ATIVA']
        sales_estornadas = [s for s in sales if s.status == 'ESTORNADA']
        total_ativas = sum(s.amount for s in sales_ativas)
        total_estornos = sum(s.amount for s in sales_estornadas)

        commission = SellerCommission.objects.filter(
            seller=seller,
            period=statement_period,
        ).select_related('period').first() if statement_period else None

        if commission and commission.period.status != CommissionPeriod.Status.ABERTA:
            comissao_valor = commission.commission_amount
            comissao_taxa = float(commission.commission_rate) * 100
            comissao_status = commission.get_status_display()
        elif statement_period:
            est, _ = calculate_estimated_commission_for_period(seller, statement_period)
            comissao_valor = est
            comissao_taxa = float(get_commission_rate(seller)) * 100
            comissao_status = 'Estimativa'
        else:
            from app.apps.commissions.services import calculate_estimated_commission
            est, _ = calculate_estimated_commission(seller, month_int, year_int)
            comissao_valor = est
            comissao_taxa = float(get_commission_rate(seller)) * 100
            comissao_status = 'Estimativa'

        from app.apps.sellers.models import SellerGoal
        goal = SellerGoal.objects.filter(
            seller=seller, month=month_int, year=year_int,
        ).first()

        semana = ['Segunda', 'Terca', 'Quarta', 'Quinta', 'Sexta', 'Sabado', 'Domingo']
        sales_data = []
        for s in sales:
            sales_data.append({
                'date': s.sale_date.strftime('%d/%m/%Y'),
                'weekday': semana[s.sale_date.weekday()],
                'amount': s.amount,
                'origin': s.get_origin_display(),
                'notes': s.notes or '',
                'status': 'Estornada' if s.status == 'ESTORNADA' else 'Ativa',
                'is_estornada': s.status == 'ESTORNADA',
            })

        if statement_period:
            competencia_label = statement_period.display_label
            periodo_range = (
                f'{start.strftime("%d/%m/%Y")} a {end.strftime("%d/%m/%Y")}'
            )
        else:
            competencia_label = f'{month_int:02d}/{year_int} (mes calendario — sem competencia registrada)'
            periodo_range = f'{start.strftime("%d/%m/%Y")} a {end.strftime("%d/%m/%Y")}'

        from django.template.loader import render_to_string
        from weasyprint import HTML

        html = render_to_string('reports/extrato_vendedor.html', {
            'logo_url': 'file://' + str(settings.BASE_DIR / 'static' / 'img' / 'vidalys-merito-logo.png'),
            'seller': seller,
            'tenant': seller.tenant,
            'competencia': competencia_label,
            'periodo_range': periodo_range,
            'is_month_fallback': is_month_fallback,
            'data_geracao': today.strftime('%d/%m/%Y'),
            'sales_data': sales_data,
            'total_ativas': total_ativas,
            'total_estornos': total_estornos,
            'comissao_valor': comissao_valor,
            'comissao_taxa': comissao_taxa,
            'comissao_status': comissao_status,
            'goal': goal,
        })

        pdf = HTML(string=html).write_pdf()

        response = HttpResponse(pdf, content_type='application/pdf')
        response['Content-Disposition'] = (
            f'attachment; filename="extrato_{seller.name}_{year_int}_{month_int:02d}.pdf"'
        )
        return response


@extend_schema(responses={(200, 'application/pdf'): OpenApiTypes.BINARY})
class MonthlyReportView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated, IsManagerOrAdmin]

    def get(self, request, year=None, month=None):
        tenant = request.user.tenant
        hoje = timezone.localdate()
        try:
            year_int = int(year) if year else hoje.year
            month_int = int(month) if month else hoje.month
        except (ValueError, TypeError):
            return Response({'error': 'Mes/ano invalidos.'}, status=400)

        from app.apps.commissions.services import calculate_estimated_commission
        period = get_period_by_legacy_label(tenant, month_int, year_int)
        sellers = Seller.objects.filter(tenant=tenant, is_active=True)
        sellers_data = []
        total_sold = 0
        total_commission_aberta = 0
        total_commission_fechada = 0
        total_commission_paga = 0

        for seller in sellers:
            est_commission, est_total = calculate_estimated_commission(seller, month_int, year_int)
            sc = SellerCommission.objects.filter(
                seller=seller,
                period__month=month_int,
                period__year=year_int,
            ).select_related('period').first()

            if sc:
                if sc.status in ('FECHADA', 'PAGA', 'AJUSTADA'):
                    st = sc
                    commission_amount = sc.commission_amount
                else:
                    commission_amount = est_commission
                    st = sc
            else:
                commission_amount = est_commission
                st = None

            status = st.get_status_display() if st else 'Estimativa'
            if st and st.status == 'FECHADA':
                total_commission_fechada += commission_amount
            elif st and st.status == 'PAGA':
                total_commission_paga += commission_amount
            else:
                total_commission_aberta += commission_amount

            sellers_data.append({
                'name': seller.name,
                'total_sold': est_total,
                'commission': commission_amount,
                'commission_rate': float(get_commission_rate(seller)) * 100,
                'status': status,
            })
            total_sold += est_total

        sellers_data.sort(key=lambda s: s['total_sold'], reverse=True)
        total_commissions = total_commission_aberta + total_commission_fechada + total_commission_paga

        prev_period = None
        if period:
            prev_period = CommissionPeriod.objects.filter(
                tenant=tenant,
                end_date__lt=period.start_date,
            ).exclude(
                status=CommissionPeriod.Status.CANCELADA,
            ).order_by('-end_date').first()

        if prev_period:
            prev_total = Sale.objects.filter(
                tenant=tenant,
                status='ATIVA',
                sale_date__gte=prev_period.start_date,
                sale_date__lte=prev_period.end_date,
            ).aggregate(t=Sum('amount'))['t'] or 0
        else:
            prev_total = 0
        variacao = round((total_sold - prev_total) / prev_total * 100) if prev_total > 0 else None

        from django.template.loader import render_to_string
        from weasyprint import HTML

        html = render_to_string('reports/relatorio_mensal.html', {
            'logo_url': 'file://' + str(settings.BASE_DIR / 'static' / 'img' / 'vidalys-merito-logo.png'),
            'tenant': tenant,
            'competencia': f'{month_int:02d}/{year_int}',
            'data_geracao': hoje.strftime('%d/%m/%Y'),
            'sellers_data': sellers_data,
            'total_sold': total_sold,
            'total_commissions': total_commissions,
            'total_aberta': total_commission_aberta,
            'total_fechada': total_commission_fechada,
            'total_paga': total_commission_paga,
            'num_sellers': sellers.count(),
            'prev_total': prev_total,
            'variacao': variacao,
        })

        pdf = HTML(string=html).write_pdf()
        response = HttpResponse(pdf, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="relatorio_{year_int}_{month_int:02d}.pdf"'
        return response


class AccountingExportView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated, IsFinancialOrAdmin | IsManagerOrAdmin]

    def get(self, request, year=None, month=None):
        tenant = request.user.tenant
        hoje = timezone.localdate()
        try:
            year_int = int(year) if year else hoje.year
            month_int = int(month) if month else hoje.month
        except (ValueError, TypeError):
            return Response({'error': 'Mes/ano invalidos.'}, status=400)

        features = getattr(settings, 'PLAN_FEATURES', {}).get(tenant.plan, {})
        if not features.get('export_contabil'):
            return Response(
                {'detail': 'Disponivel a partir do plano Pro. Faca upgrade em Configuracoes > Assinatura.',
                 'upgrade_required': True},
                status=403,
            )

        from app.apps.commissions.exports import build_accounting_zip

        zip_bytes = build_accounting_zip(tenant, month_int, year_int)

        log_action(request, 'accounting_export', changes={
            'month': month_int, 'year': year_int, 'tenant': str(tenant.uuid),
        })

        resp = HttpResponse(zip_bytes, content_type='application/zip')
        resp['Content-Disposition'] = f'attachment; filename="contabilidade_{year_int}_{month_int:02d}.zip"'
        return resp


class AccountingEmailView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated, IsFinancialOrAdmin | IsManagerOrAdmin]

    def post(self, request, year=None, month=None):
        tenant = request.user.tenant
        features = getattr(settings, 'PLAN_FEATURES', {}).get(tenant.plan, {})
        if not features.get('export_contabil'):
            return Response(
                {'detail': 'Disponivel a partir do plano Pro.', 'upgrade_required': True},
                status=403,
            )
        if not tenant.accountant_email:
            return Response(
                {'detail': 'Cadastre o e-mail do contador em Configuracoes.'},
                status=400,
            )
        try:
            year_int = int(year)
            month_int = int(month)
        except (ValueError, TypeError):
            return Response({'error': 'Mes/ano invalidos.'}, status=400)

        from app.apps.notifications.tasks import send_accounting_package_email
        send_accounting_package_email.delay(
            str(tenant.uuid), month_int, year_int,
            requested_by_user_id=str(request.user.pk),
        )
        return Response({'detail': f'Envio agendado para {tenant.accountant_email}.'}, status=202)


class SellerDayJustificationViewSet(viewsets.ModelViewSet):
    """CRUD de gestor para justificativas de dia sem lancamento (Prompt 47).

    Fundacao de backend: NAO integra com fechamento/comissao/ranking/mobile.
    Permissao: gestao operacional (MANAGER/ADMIN). SELLER e FINANCEIRO -> 403.
    Todo queryset e limitado ao tenant do usuario (uuid de outro tenant -> 404).
    """

    permission_classes = [IsAuthenticated, IsManagerOrAdmin]

    def get_serializer_class(self):
        if self.action == 'create':
            return SellerDayJustificationCreateSerializer
        if self.action in ('update', 'partial_update'):
            return SellerDayJustificationUpdateSerializer
        return SellerDayJustificationSerializer

    def get_queryset(self):
        from app.apps.sellers.models import SellerDayJustification
        user = self.request.user
        qs = SellerDayJustification.objects.select_related(
            'seller', 'created_by', 'updated_by',
        ).filter(tenant=user.tenant)

        params = self.request.query_params
        seller = params.get('seller')
        if seller:
            qs = qs.filter(seller__uuid=seller)
        single_date = params.get('date')
        if single_date:
            qs = qs.filter(date=single_date)
        start_date = params.get('start_date')
        if start_date:
            qs = qs.filter(date__gte=start_date)
        end_date = params.get('end_date')
        if end_date:
            qs = qs.filter(date__lte=end_date)
        reason = params.get('reason')
        if reason:
            qs = qs.filter(reason=reason)
        return qs.order_by('-date', '-created_at')

    def perform_destroy(self, instance):
        from app.apps.sellers.services import delete_day_justification
        delete_day_justification(justification=instance, user=self.request.user)

    # --- Mapeamento de erros de dominio para status HTTP (LOTE 2/3) ---
    def _map_justification_errors(self, func):
        from app.apps.sellers.services import (
            JustificationConflictError, JustificationLockedError,
            JustificationError,
        )
        try:
            return func()
        except (JustificationConflictError, JustificationLockedError) as exc:
            return Response({'detail': str(exc)}, status=409)
        except JustificationError as exc:
            return Response({'detail': str(exc)}, status=400)

    def create(self, request, *args, **kwargs):
        return self._map_justification_errors(
            lambda: super(SellerDayJustificationViewSet, self).create(
                request, *args, **kwargs,
            )
        )

    def update(self, request, *args, **kwargs):
        return self._map_justification_errors(
            lambda: super(SellerDayJustificationViewSet, self).update(
                request, *args, **kwargs,
            )
        )

    def destroy(self, request, *args, **kwargs):
        return self._map_justification_errors(
            lambda: super(SellerDayJustificationViewSet, self).destroy(
                request, *args, **kwargs,
            )
        )

    @action(
        detail=True, methods=['post'],
        permission_classes=[IsAuthenticated, IsManagerOrAdmin],
        url_path='replace-with-sale',
    )
    def replace_with_sale(self, request, pk=None):
        """Substitui explicitamente a justificativa por uma venda (LOTE 2)."""
        from app.apps.sellers.services import (
            replace_justification_with_sale,
            JustificationConflictError, JustificationLockedError,
            JustificationError,
        )
        justification = self.get_object()
        amount = request.data.get('amount')
        notes = request.data.get('notes', '')
        try:
            sale = replace_justification_with_sale(
                justification=justification, amount=amount,
                user=request.user, notes=notes,
            )
        except (JustificationConflictError, JustificationLockedError) as exc:
            return Response({'detail': str(exc)}, status=409)
        except JustificationError as exc:
            return Response({'detail': str(exc)}, status=400)
        return Response(
            {'sale_uuid': str(sale.uuid), 'amount': sale.amount,
             'sale_date': sale.sale_date.isoformat()},
            status=201,
        )


class SellerDayStatusView(generics.GenericAPIView):
    """Estados operacionais dos dias de um vendedor numa competencia (LOTE 7).

    Manager/Admin. Tenant-scoped: seller/period de outro tenant -> 404.
    """

    permission_classes = [IsAuthenticated, IsManagerOrAdmin]

    def get(self, request, seller_id=None):
        from app.apps.commissions.day_status import get_period_day_statuses
        tenant = request.user.tenant
        try:
            seller = Seller.objects.get(uuid=seller_id, tenant=tenant)
        except Seller.DoesNotExist:
            from django.http import Http404
            raise Http404('Vendedor nao encontrado.')

        period_uuid = request.query_params.get('period')
        if not period_uuid:
            return Response({'detail': 'period e obrigatorio.'}, status=400)
        try:
            period = CommissionPeriod.objects.get(
                uuid=period_uuid, tenant=tenant,
            )
        except (CommissionPeriod.DoesNotExist, ValueError, ValidationError):
            from django.http import Http404
            raise Http404('Competencia nao encontrada.')

        result = get_period_day_statuses(tenant, seller, period)
        days = [
            {
                'date': d['date'].isoformat(),
                'is_expected_day': d['is_expected_day'],
                'status': d['status'],
                'active_sales_count': d['active_sales_count'],
                'active_sales_total': d['active_sales_total'],
                'justification_uuid': d['justification_uuid'],
                'justification_reason': d['justification_reason'],
                'justification_reason_display': d['justification_reason_display'],
                'justification_notes': d['justification_notes'],
                'can_manage_justification': d['can_manage_justification'],
            }
            for d in result['days']
        ]
        return Response({'summary': result['summary'], 'days': days})
