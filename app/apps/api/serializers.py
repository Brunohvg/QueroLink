from datetime import date
import logging

from rest_framework import serializers
from django.db import IntegrityError, transaction
from django.utils import timezone

from app.apps.sales.models import Sale
from app.apps.sellers.models import Seller
from app.apps.sellers.capacity import SellerCapacityExceeded, ensure_seller_capacity
from app.apps.commissions.models import (
    CommissionPeriod,
    SellerCommission,
    CommissionAdjustment,
)
from app.apps.accounts.models import Tenant, User
from app.apps.accounts.validators import clean_phone, validate_phone_br
from app.apps.accounts.fields import compute_hash

logger = logging.getLogger(__name__)


def _create_unique_seller_user(tenant, seller_name, used_usernames=None):
    used_usernames = used_usernames if used_usernames is not None else set()
    base = slugify(seller_name) or 'seller'
    n = 1

    for _attempt in range(10):
        username = base if n == 1 else f'{base}-{n}'
        while username in used_usernames or User.objects.filter(username=username).exists():
            n += 1
            username = f'{base}-{n}'

        password = get_random_string(12)
        try:
            with transaction.atomic():
                user = User.objects.create_user(
                    username=username,
                    password=password,
                    role=User.Role.SELLER,
                    tenant=tenant,
                )
        except IntegrityError:
            used_usernames.add(username)
            n += 1
            continue

        used_usernames.add(username)
        return user, username, password

    raise serializers.ValidationError({
        'detail': 'Nao foi possivel gerar um username unico para o vendedor.',
    })


class SellerSerializer(serializers.ModelSerializer):
    username = serializers.CharField(source='user.username', read_only=True)
    cpf_formatted = serializers.CharField(read_only=True)

    class Meta:
        model = Seller
        fields = [
            'uuid', 'name', 'phone', 'cpf', 'cpf_formatted', 'commission_rate',
            'is_active', 'username', 'created_at',
        ]
        read_only_fields = ['uuid', 'username', 'created_at', 'cpf_formatted']

    def validate_phone(self, value):
        cleaned = clean_phone(value)
        if not validate_phone_br(cleaned):
            raise serializers.ValidationError(
                'Telefone invalido. Informe um numero com DDD (10 ou 11 digitos).'
            )
        tenant = self.context['request'].user.tenant
        qs = Seller.objects.filter(tenant=tenant, phone_hash=compute_hash(cleaned))
        if self.instance:
            qs = qs.exclude(pk=self.instance.pk)
        if qs.exists():
            raise serializers.ValidationError(
                'Este telefone ja esta cadastrado para outro vendedor.'
            )
        return cleaned

    def validate_cpf(self, value):
        if not value:
            return None
        from app.apps.sellers.validators import normalize_and_validate_cpf
        try:
            digits = normalize_and_validate_cpf(value)
        except ValueError as e:
            raise serializers.ValidationError(str(e))
        tenant = self.context['request'].user.tenant
        qs = Seller.objects.filter(tenant=tenant, cpf=digits)
        if self.instance:
            qs = qs.exclude(pk=self.instance.pk)
        if qs.exists():
            raise serializers.ValidationError('Ja existe um vendedor com este CPF.')
        return digits


class SellerCreateSerializer(serializers.Serializer):
    name = serializers.CharField(max_length=100)
    phone = serializers.CharField(max_length=20)
    cpf = serializers.CharField(max_length=14, required=False, allow_blank=True)

    def validate_phone(self, value):
        cleaned = clean_phone(value)
        if not validate_phone_br(cleaned):
            raise serializers.ValidationError(
                'Telefone invalido. Informe um numero com DDD (10 ou 11 digitos).'
            )
        request = self.context.get('request')
        if request and request.user.tenant:
            if Seller.objects.filter(
                tenant=request.user.tenant, phone_hash=compute_hash(cleaned)
            ).exists():
                raise serializers.ValidationError(
                    'Este telefone ja esta cadastrado para outro vendedor.'
                )
        return cleaned

    def validate_cpf(self, value):
        if not value:
            return None
        from app.apps.sellers.validators import normalize_and_validate_cpf
        try:
            digits = normalize_and_validate_cpf(value)
        except ValueError as e:
            raise serializers.ValidationError(str(e))
        request = self.context.get('request')
        if request and request.user.tenant:
            if Seller.objects.filter(
                tenant=request.user.tenant, cpf=digits
            ).exists():
                raise serializers.ValidationError('Ja existe um vendedor com este CPF.')
        return digits

    def create(self, validated_data):
        request = self.context['request']
        tenant = request.user.tenant

        try:
            with transaction.atomic():
                locked_tenant = Tenant.objects.select_for_update().get(pk=tenant.pk)
                ensure_seller_capacity(locked_tenant, requested=1)
                user, username, password = _create_unique_seller_user(
                    locked_tenant,
                    validated_data['name'],
                )

                seller = Seller.objects.create(
                    tenant=locked_tenant,
                    user=user,
                    name=validated_data['name'],
                    phone=validated_data['phone'],
                    cpf=validated_data.get('cpf', '') or None,
                    commission_rate=locked_tenant.default_commission_rate,
                )
        except SellerCapacityExceeded as e:
            raise serializers.ValidationError({'detail': str(e)})

        try:
            from app.apps.notifications.tasks import notify_seller_credentials
            notify_seller_credentials(seller, password)
        except Exception:
            logger.error(
                'Failed to send credentials notification for seller %s',
                seller.uuid, exc_info=True
            )

        return {
            'uuid': str(seller.uuid),
            'name': seller.name,
            'phone': seller.phone,
            'username': username,
            'message': 'Vendedor criado com sucesso. As credenciais serao enviadas via WhatsApp.',
        }


class SellerImportSerializer(serializers.Serializer):
    file = serializers.FileField()

    def validate_file(self, value):
        name = value.name.lower()
        if not (name.endswith('.csv') or name.endswith('.xlsx')):
            raise serializers.ValidationError(
                'Formato de arquivo invalido. Envie um arquivo .csv ou .xlsx.'
            )
        if value.size > 5 * 1024 * 1024:
            raise serializers.ValidationError(
                'Arquivo muito grande. Tamanho maximo permitido: 5MB.'
            )
        if name.endswith('.xlsx'):
            header = value.read(512)
            value.seek(0)
            if not header.startswith(b'PK\x03\x04'):
                raise serializers.ValidationError(
                    'Arquivo .xlsx invalido ou corrompido.'
                )
        return value

    def _parse_file(self, file):
        import csv
        import io

        content = file.read()

        if file.name.lower().endswith('.csv'):
            for encoding in ('utf-8-sig', 'latin-1', 'cp1252'):
                try:
                    decoded = content.decode(encoding)
                    break
                except (UnicodeDecodeError, LookupError):
                    continue
            else:
                raise serializers.ValidationError(
                    'Nao foi possivel decodificar o arquivo. Tente salvar como UTF-8.'
                )

            rows = self._parse_csv(decoded)
            if not rows:
                raise serializers.ValidationError(
                    'Nenhuma linha encontrada. Verifique se o arquivo tem cabecalho '
                    'com as colunas "nome" e "telefone".'
                )
            return rows

        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        ws = wb.active
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        return [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]

    def _parse_csv(self, decoded):
        import csv
        import io
        import logging

        logger = logging.getLogger(__name__)
        delimiters = [',', ';', '\t', '|']

        for delim in delimiters:
            reader = csv.DictReader(io.StringIO(decoded), delimiter=delim)
            try:
                rows = [row for row in reader]
            except (csv.Error, UnicodeDecodeError) as e:
                logger.debug("CSV parsing failed with delimiter %r: %s", delim, e)
                continue
            if not rows:
                continue
            keys = set(rows[0].keys())
            if 'nome' in keys or 'name' in keys or 'telefone' in keys or 'phone' in keys:
                return rows

        reader = csv.DictReader(io.StringIO(decoded))
        return [row for row in reader]

    def _build_import_plan(self, rows, tenant):
        errors = []
        pending = []
        used_phones = set(
            Seller.objects.filter(tenant=tenant)
            .values_list('phone_hash', flat=True)
        )
        used_cpfs = set(
            Seller.objects.filter(tenant=tenant, cpf__isnull=False)
            .values_list('cpf', flat=True)
        )
        batch_phones = set()
        batch_cpfs = set()

        for i, row in enumerate(rows, start=2):
            name = (row.get('nome') or row.get('name') or '').strip()
            phone = (row.get('telefone') or row.get('phone') or '').strip()

            if not name:
                errors.append({'linha': i, 'erro': 'Nome nao informado.'})
                continue
            if not phone:
                errors.append({'linha': i, 'erro': 'Telefone nao informado.'})
                continue

            cleaned_phone = clean_phone(phone)
            if not validate_phone_br(cleaned_phone):
                errors.append({'linha': i, 'erro': f'Telefone invalido: {phone}'})
                continue

            phone_hash = compute_hash(cleaned_phone)
            if phone_hash in used_phones or phone_hash in batch_phones:
                errors.append({'linha': i, 'erro': f'Telefone ja cadastrado: {phone}'})
                continue

            cpf_raw = (row.get('cpf') or row.get('CPF') or '').strip()
            cpf_value = None
            if cpf_raw:
                from app.apps.sellers.validators import normalize_and_validate_cpf
                try:
                    candidate = normalize_and_validate_cpf(cpf_raw)
                    if candidate in used_cpfs or candidate in batch_cpfs:
                        errors.append({
                            'linha': i,
                            'erro': f'CPF ignorado: duplicado (ja cadastrado): {cpf_raw}',
                        })
                    else:
                        cpf_value = candidate
                        batch_cpfs.add(candidate)
                except ValueError as e:
                    errors.append({
                        'linha': i,
                        'erro': f'CPF ignorado: {str(e).lower()}: {cpf_raw}',
                    })

            batch_phones.add(phone_hash)
            pending.append({
                'line': i,
                'name': name,
                'phone': phone,
                'cleaned_phone': cleaned_phone,
                'cpf': cpf_value,
            })

        return pending, errors

    def create(self, validated_data):
        file = validated_data['file']
        rows = self._parse_file(file)

        request = self.context['request']
        tenant = request.user.tenant

        created = []
        credentials = []
        used_usernames = set()

        try:
            with transaction.atomic():
                locked_tenant = Tenant.objects.select_for_update().get(pk=tenant.pk)
                pending, errors = self._build_import_plan(rows, locked_tenant)
                ensure_seller_capacity(locked_tenant, requested=len(pending))

                for item in pending:
                    user, username, password = _create_unique_seller_user(
                        locked_tenant,
                        item['name'],
                        used_usernames,
                    )
                    seller = Seller.objects.create(
                        tenant=locked_tenant,
                        user=user,
                        name=item['name'],
                        phone=item['cleaned_phone'],
                        commission_rate=locked_tenant.default_commission_rate,
                        cpf=item['cpf'],
                    )
                    used_usernames.add(username)
                    credentials.append((seller, password))
                    created.append({
                        'username': username,
                        'name': item['name'],
                        'phone': item['phone'],
                        'uuid': str(seller.uuid),
                    })
        except SellerCapacityExceeded as e:
            raise serializers.ValidationError({'detail': str(e)})

        for seller, password in credentials:
            try:
                from app.apps.notifications.tasks import notify_seller_credentials
                notify_seller_credentials(seller, password)
            except Exception:
                logger.error(
                    'Failed to send credentials notification for seller %s',
                    seller.uuid, exc_info=True
                )

        return {
            'total': len(rows),
            'criados': len(created),
            'erros': len(errors),
            'created': created,
            'errors': errors,
        }


class SaleSerializer(serializers.ModelSerializer):
    seller_uuid = serializers.CharField(source='seller.uuid', read_only=True)
    seller_name = serializers.CharField(source='seller.name', read_only=True)

    class Meta:
        model = Sale
        fields = [
            'uuid', 'seller_uuid', 'seller_name', 'origin', 'amount',
            'sale_date', 'notes', 'created_at', 'updated_at',
        ]
        read_only_fields = [
            'uuid', 'seller_uuid', 'seller_name', 'created_at', 'updated_at',
        ]


class SaleCreateSerializer(serializers.ModelSerializer):
    seller = serializers.PrimaryKeyRelatedField(
        queryset=Seller.objects.all(), required=False,
    )

    class Meta:
        model = Sale
        fields = ['seller', 'origin', 'amount', 'sale_date', 'notes']

    def get_validators(self):
        return []

    def validate(self, attrs):
        request = self.context.get('request')
        user = request.user
        today = timezone.localdate()

        if user.role == User.Role.SELLER:
            attrs['seller'] = user.seller_profile

        seller = attrs.get('seller')
        if not seller:
            raise serializers.ValidationError({'seller': 'Vendedor e obrigatorio.'})
        if seller.tenant_id != user.tenant_id:
            raise serializers.ValidationError(
                {'seller': 'O vendedor nao pertence ao tenant.'}
            )

        sale_date = attrs.get('sale_date')
        if not sale_date:
            raise serializers.ValidationError(
                {'sale_date': 'Data do lancamento e obrigatoria.'}
            )

        if sale_date > today:
            raise serializers.ValidationError({
                'sale_date': 'Nao e possivel lancar vendas em data futura.'
            })

        amount = attrs.get('amount', 0)
        if amount and int(amount) > 10_000_000:
            raise serializers.ValidationError({
                'amount': 'Valor maximo e R$ 100.000,00.'
            })

        origin = attrs.get('origin', Sale.Origin.MANUAL)

        if origin == Sale.Origin.MANUAL:
            from app.apps.commissions.services import validate_sale_can_be_changed
            can_change, error_msg = validate_sale_can_be_changed(seller, sale_date, user)
            if not can_change:
                raise serializers.ValidationError({'sale_date': error_msg})

            existing = Sale.objects.filter(
                seller=seller,
                sale_date=sale_date,
                origin=Sale.Origin.MANUAL,
            )
            if self.instance:
                existing = existing.exclude(pk=self.instance.pk)
            if existing.exists():
                raise serializers.ValidationError({
                    'sale_date': (
                        'Ja existe um lancamento manual para este vendedor '
                        'nesta data. Para alterar, edite o lancamento existente.'
                    ),
                })

        return attrs

    def create(self, validated_data):
        request = self.context.get('request')
        validated_data['tenant'] = request.user.tenant
        validated_data['created_by'] = request.user
        return super().create(validated_data)

    def update(self, instance, validated_data):
        request = self.context.get('request')
        validated_data['updated_by'] = request.user
        return super().update(instance, validated_data)


class CommissionAdjustmentSerializer(serializers.ModelSerializer):
    adjusted_by_name = serializers.CharField(
        source='adjusted_by.username', read_only=True,
    )

    class Meta:
        model = CommissionAdjustment
        fields = [
            'id', 'previous_amount', 'new_amount', 'difference',
            'reason', 'adjusted_by_name', 'created_at',
        ]
        read_only_fields = [
            'id', 'previous_amount', 'new_amount', 'difference',
            'adjusted_by_name', 'created_at',
        ]


class SellerCommissionReadSerializer(serializers.ModelSerializer):
    seller_name = serializers.CharField(source='seller.name', read_only=True)
    seller_uuid = serializers.CharField(source='seller.uuid', read_only=True)
    period_status = serializers.CharField(
        source='period.status', read_only=True,
    )
    adjustments = CommissionAdjustmentSerializer(many=True, read_only=True)
    is_editable = serializers.BooleanField(read_only=True)

    class Meta:
        model = SellerCommission
        fields = [
            'id', 'seller_uuid', 'seller_name',
            'status', 'operational_status',
            'total_sold_amount', 'commission_rate', 'commission_amount',
            'expected_working_days', 'submitted_days_count', 'missing_days_count',
            'frozen_total_sold_amount', 'frozen_commission_rate',
            'frozen_commission_amount',
            'period_status',
            'closed_at', 'reopened_at', 'reopen_reason',
            'paid_at', 'paid_amount', 'payment_date',
            'payment_method', 'payment_notes',
            'adjustments', 'is_editable',
        ]
        read_only_fields = fields


class CommissionPeriodSerializer(serializers.ModelSerializer):
    seller_commissions = serializers.SerializerMethodField()
    is_current_month = serializers.SerializerMethodField()

    class Meta:
        model = CommissionPeriod
        fields = [
            'uuid', 'tenant', 'month', 'year', 'status',
            'expected_working_days', 'notes',
            'closed_at', 'paid_at', 'adjusted_at',
            'adjustment_reason', 'cancelled_at', 'cancel_reason',
            'sent_to_accounting_at',
            'created_at', 'seller_commissions', 'is_current_month',
        ]
        read_only_fields = [
            'uuid', 'tenant', 'closed_at', 'paid_at', 'adjusted_at',
            'adjustment_reason', 'cancelled_at', 'cancel_reason',
            'created_at', 'seller_commissions', 'is_current_month',
        ]

    def get_seller_commissions(self, obj):
        commissions = obj.seller_commissions.select_related('seller').all()
        import copy
        result = []
        for sc in commissions:
            if sc.is_editable:
                temp = copy.copy(sc)
                temp.recalculate(commit=False)
                result.append(temp)
            else:
                result.append(sc)
        return SellerCommissionReadSerializer(result, many=True).data

    def get_is_current_month(self, obj):
        hoje = timezone.localdate()
        return obj.month == hoje.month and obj.year == hoje.year


class CommissionPeriodCreateSerializer(serializers.ModelSerializer):
    class Meta:
        model = CommissionPeriod
        fields = ['month', 'year', 'expected_working_days', 'notes']
        extra_kwargs = {
            'expected_working_days': {'required': False},
            'notes': {'required': False},
        }

    def validate_month(self, value):
        if value < 1 or value > 12:
            raise serializers.ValidationError('Mes deve estar entre 1 e 12.')
        return value

    def validate_year(self, value):
        if value < 2000 or value > 2100:
            raise serializers.ValidationError('Ano invalido.')
        return value

    def validate(self, attrs):
        request = self.context.get('request')
        tenant = request.user.tenant
        month = attrs.get('month')
        year = attrs.get('year')
        if CommissionPeriod.objects.filter(
            tenant=tenant, month=month, year=year,
        ).exists():
            raise serializers.ValidationError(
                f'A competencia {month:02d}/{year} ja existe. '
                'Use Atualizar valores para sincronizar os dados.'
            )
        return attrs


class ChangePasswordSerializer(serializers.Serializer):
    current_password = serializers.CharField(required=True)
    new_password = serializers.CharField(required=True, min_length=8)

    def validate_current_password(self, value):
        user = self.context['request'].user
        if not user.check_password(value):
            raise serializers.ValidationError('Senha atual incorreta.')
        return value

    def save(self):
        user = self.context['request'].user
        user.set_password(self.validated_data['new_password'])
        user.save()
        return user


def slugify(value):
    from django.utils.text import slugify as _slugify
    return _slugify(value)


def get_random_string(length):
    from django.utils.crypto import get_random_string as _get_random_string
    return _get_random_string(length)
