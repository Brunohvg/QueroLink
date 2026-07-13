from datetime import date
from decimal import Decimal
import io
from unittest.mock import patch

from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.core.cache import cache
from django.db import connection
from rest_framework.test import APIClient

from app.apps.accounts.models import Tenant, User
from app.apps.sellers.models import Seller, SellerDayJustification
from app.apps.commissions.models import CommissionPeriod, SellerCommission
from app.apps.sales.models import Sale, SaleChangeLog


class CompetenciaFirstTest(TestCase):
    """PROMPT_42 - competencia-first: uma tela = um regime."""

    def setUp(self):
        cache.clear()
        self.tenant = Tenant.objects.create(
            company_name='Comp First', slug='comp-first',
            default_commission_rate=Decimal('0.01'),
            period_start_day=21,
        )
        self.tenant2 = Tenant.objects.create(
            company_name='Outro Tenant', slug='outro-tenant',
            default_commission_rate=Decimal('0.01'),
        )
        self.manager = User.objects.create_user(
            username='gestor', password='gestor123',
            role=User.Role.MANAGER, tenant=self.tenant,
        )
        self.manager2 = User.objects.create_user(
            username='gestor2', password='gestor123',
            role=User.Role.MANAGER, tenant=self.tenant2,
        )
        self.seller_user = User.objects.create_user(
            username='vendedor', password='senha123',
            role=User.Role.SELLER, tenant=self.tenant,
        )
        self.seller = Seller.objects.create(
            tenant=self.tenant, name='Vendedor Teste', phone='11999998888',
            user=self.seller_user, commission_rate=Decimal('0.01'),
        )
        # Competencia 21/06 a 20/07 (ciclo dinamico)
        self.period = CommissionPeriod.objects.create(
            tenant=self.tenant, month=7, year=2026,
            start_date=date(2026, 6, 21), end_date=date(2026, 7, 20),
            label='Julho/2026',
        )
        # Vendas dentro do range: 25/06 e 05/07
        Sale.objects.create(
            tenant=self.tenant, seller=self.seller, origin=Sale.Origin.MANUAL,
            amount=5227922, sale_date='2026-06-25', created_by=self.seller_user,
        )
        Sale.objects.create(
            tenant=self.tenant, seller=self.seller, origin=Sale.Origin.MANUAL,
            amount=4723937, sale_date='2026-07-05', created_by=self.seller_user,
        )
        # Venda fora do range (21/07)
        Sale.objects.create(
            tenant=self.tenant, seller=self.seller, origin=Sale.Origin.MANUAL,
            amount=1000000, sale_date='2026-07-21', created_by=self.seller_user,
        )

    def _auth(self, user, password):
        client = APIClient()
        resp = client.post(reverse('api-login'), {
            'username': user.username, 'password': password,
        }, format='json')
        client.credentials(HTTP_AUTHORIZATION=f'Bearer {resp.data["access"]}')
        return client

    # 1 + 2: vendas 25/06 e 05/07 entram juntas; 21/07 fica de fora
    def test_period_range_aggregates_both_months(self):
        from app.apps.commissions.services import get_manual_sales_total_for_period
        total = get_manual_sales_total_for_period(self.seller, self.period)
        self.assertEqual(total, 5227922 + 4723937)

    def test_sale_outside_range_excluded(self):
        from app.apps.commissions.services import get_manual_sales_total_for_period
        total = get_manual_sales_total_for_period(self.seller, self.period)
        self.assertNotIn(1000000, [total])
        self.assertEqual(total, 9951859)

    # 4 + 5 + 3: detalhe do vendedor usa o mesmo range; card usa periodo selecionado
    def test_seller_detail_period_total_and_commission_same_range(self):
        client = self._auth(self.manager, 'gestor123')
        url = f'/api/manager/seller/{self.seller.uuid}/'
        resp = client.get(url + f'?period={self.period.uuid}')
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertEqual(data['manual_total'], 9951859)
        spc = data['selected_period_commission']
        self.assertIsNotNone(spc)
        # comissao == taxa x total (1%)
        self.assertEqual(spc['total_sold_amount'], 9951859)
        self.assertEqual(spc['commission_amount'], round(9951859 * 0.01))
        self.assertEqual(spc['commission_amount'], 99519)

    # 6: home gestor soma junho + julho pelo range
    def test_dashboard_summary_sums_range(self):
        client = self._auth(self.manager, 'gestor123')
        url = reverse('api-dashboard-summary')
        resp = client.get(url + f'?period={self.period.uuid}')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.json()['total_vendido'], 9951859)

    # 16: UUID de outro tenant retorna 404
    def test_period_from_other_tenant_returns_404(self):
        other_period = CommissionPeriod.objects.create(
            tenant=self.tenant2, month=7, year=2026,
            start_date=date(2026, 6, 21), end_date=date(2026, 7, 20),
        )
        client = self._auth(self.manager, 'gestor123')
        url = f'/api/manager/seller/{self.seller.uuid}/'
        resp = client.get(url + f'?period={other_period.uuid}')
        self.assertEqual(resp.status_code, 404)

    # 17: competencia cancelada nao aparece no seletor
    def test_cancelled_period_not_in_selector(self):
        from app.apps.commissions.services import get_selectable_periods
        cancelled = CommissionPeriod.objects.create(
            tenant=self.tenant, month=5, year=2026,
            start_date=date(2026, 4, 21), end_date=date(2026, 5, 20),
            status=CommissionPeriod.Status.CANCELADA,
        )
        uuids = [str(p.uuid) for p in get_selectable_periods(self.tenant)]
        self.assertNotIn(str(cancelled.uuid), uuids)
        self.assertIn(str(self.period.uuid), uuids)

    # 18: competencia fechada usa valor congelado
    def test_closed_period_uses_frozen_value(self):
        self.period.status = CommissionPeriod.Status.FECHADA
        self.period.save()
        sc = SellerCommission.objects.create(
            period=self.period, seller=self.seller,
            commission_rate=Decimal('0.01'),
            status=SellerCommission.Status.FECHADA,
            total_sold_amount=9951859, commission_amount=88888,
        )
        client = self._auth(self.manager, 'gestor123')
        url = f'/api/manager/seller/{self.seller.uuid}/'
        resp = client.get(url + f'?period={self.period.uuid}')
        spc = resp.json()['selected_period_commission']
        self.assertEqual(spc['commission_amount'], 88888)
        self.assertFalse(spc['is_estimated'])

        summary = client.get(
            reverse('api-seller-dashboard-summary')
            + f'?period={self.period.uuid}',
        ).json()['sellers'][0]
        self.assertEqual(summary['commission_amount'], 88888)

    # 19: periodo customizado nao exibe comissao oficial
    def test_custom_range_no_selected_commission(self):
        client = self._auth(self.manager, 'gestor123')
        url = f'/api/manager/seller/{self.seller.uuid}/'
        resp = client.get(url + '?start=2026-06-25&end=2026-07-05')
        data = resp.json()
        self.assertTrue(data['is_custom_range'])
        self.assertIsNone(data['selected_period_commission'])

    # 11: comparacao entre competencias - venda 25/06 no bucket de Julho
    def test_comparison_bucket_by_period_range(self):
        client = self._auth(self.manager, 'gestor123')
        url = f'/api/manager/seller/{self.seller.uuid}/'
        resp = client.get(url + f'?period={self.period.uuid}')
        comparison = resp.json()['comparison']
        july = [c for c in comparison if c['month'] == 'Julho/2026']
        self.assertEqual(len(july), 1)
        self.assertEqual(july[0]['total'], 9951859)

    # 10: CSV/XLSX/PDF com period retornam as mesmas vendas
    def test_exports_with_period_same_sales(self):
        client = self._auth(self.manager, 'gestor123')
        base = f'/api/manager/seller/{self.seller.uuid}/csv/'
        resp = client.get(base + f'?period={self.period.uuid}')
        self.assertEqual(resp.status_code, 200)
        content = resp.content.decode('utf-8')
        self.assertIn('2026-06-25', content)
        self.assertIn('2026-07-05', content)
        self.assertNotIn('2026-07-21', content)

    # 3: lista de vendedores - total e comissao usam o mesmo periodo (ABERTA => taxa x total)
    def test_seller_list_total_times_rate_equals_commission(self):
        client = self._auth(self.manager, 'gestor123')
        list_response = client.get(reverse('api-seller-list'))
        self.assertEqual(list_response.status_code, 200)
        self.assertEqual(
            [item['uuid'] for item in list_response.json()],
            [str(self.seller.uuid)],
        )
        detail_url = f'/api/manager/seller/{self.seller.uuid}/'
        data = client.get(
            detail_url + f'?period={self.period.uuid}',
        ).json()
        total = data['manual_total']
        commission = data['selected_period_commission']['commission_amount']
        self.assertEqual(commission, round(total * 0.01))

    def test_seller_dashboard_summary_aggregates_in_one_response(self):
        client = self._auth(self.manager, 'gestor123')
        response = client.get(
            reverse('api-seller-dashboard-summary')
            + f'?period={self.period.uuid}',
        )
        self.assertEqual(response.status_code, 200)
        sellers = response.json()['sellers']
        self.assertEqual(len(sellers), 1)
        self.assertEqual(sellers[0]['month_total'], 9951859)
        self.assertEqual(sellers[0]['commission_amount'], 99519)
        self.assertEqual(sellers[0]['last_sale_date'], '2026-07-05')
        self.assertEqual(sellers[0]['financial_status'], 'SEM_COMISSAO')
        self.assertFalse(sellers[0]['has_sale_today'])
        self.assertFalse(sellers[0]['has_justification_today'])
        self.assertFalse(sellers[0]['has_day_resolved_today'])

    def test_summary_without_commission_is_informational_when_locked(self):
        client = self._auth(self.manager, 'gestor123')
        url = (
            reverse('api-seller-dashboard-summary')
            + f'?period={self.period.uuid}'
        )
        for period_status in (
            CommissionPeriod.Status.FECHADA,
            CommissionPeriod.Status.PAGA,
        ):
            self.period.status = period_status
            self.period.save(update_fields=['status'])
            seller = client.get(url).json()['sellers'][0]
            self.assertEqual(seller['financial_status'], 'SEM_COMISSAO')
            self.assertNotEqual(seller['financial_status'], 'ABERTA')
            self.assertEqual(seller['commission_amount'], 99519)

    def test_cancelled_period_is_not_exposed_as_open_commission(self):
        self.period.status = CommissionPeriod.Status.CANCELADA
        self.period.save(update_fields=['status'])
        client = self._auth(self.manager, 'gestor123')
        response = client.get(
            reverse('api-seller-dashboard-summary')
            + f'?period={self.period.uuid}',
        )
        self.assertEqual(response.status_code, 404)
        self.assertNotContains(response, 'ABERTA', status_code=404)

    def test_existing_seller_commission_keeps_its_status(self):
        SellerCommission.objects.create(
            period=self.period,
            seller=self.seller,
            status=SellerCommission.Status.FECHADA,
            commission_rate=Decimal('0.01'),
            total_sold_amount=9951859,
            commission_amount=99519,
        )
        client = self._auth(self.manager, 'gestor123')
        seller = client.get(
            reverse('api-seller-dashboard-summary')
            + f'?period={self.period.uuid}',
        ).json()['sellers'][0]
        self.assertEqual(seller['financial_status'], 'FECHADA')
        self.assertEqual(seller['commission_amount'], 99519)

    @patch('app.apps.api.views.timezone.localdate')
    def test_active_manual_sale_resolves_today(self, localdate_mock):
        localdate_mock.return_value = date(2026, 7, 10)
        Sale.objects.create(
            tenant=self.tenant,
            seller=self.seller,
            origin=Sale.Origin.MANUAL,
            status='ATIVA',
            amount=10000,
            sale_date=date(2026, 7, 10),
            created_by=self.seller_user,
        )
        client = self._auth(self.manager, 'gestor123')
        seller = client.get(
            reverse('api-seller-dashboard-summary')
            + f'?period={self.period.uuid}',
        ).json()['sellers'][0]
        self.assertTrue(seller['has_sale_today'])
        self.assertFalse(seller['has_justification_today'])
        self.assertTrue(seller['has_day_resolved_today'])

    @patch('app.apps.api.views.timezone.localdate')
    def test_justification_resolves_today(self, localdate_mock):
        localdate_mock.return_value = date(2026, 7, 10)
        SellerDayJustification.objects.create(
            tenant=self.tenant,
            seller=self.seller,
            date=date(2026, 7, 10),
            reason=SellerDayJustification.Reason.FALTA,
            created_by=self.manager,
        )
        client = self._auth(self.manager, 'gestor123')
        seller = client.get(
            reverse('api-seller-dashboard-summary')
            + f'?period={self.period.uuid}',
        ).json()['sellers'][0]
        self.assertFalse(seller['has_sale_today'])
        self.assertTrue(seller['has_justification_today'])
        self.assertTrue(seller['has_day_resolved_today'])

    @patch('app.apps.api.views.timezone.localdate')
    def test_refunded_sale_does_not_resolve_today(self, localdate_mock):
        localdate_mock.return_value = date(2026, 7, 10)
        Sale.objects.create(
            tenant=self.tenant,
            seller=self.seller,
            origin=Sale.Origin.MANUAL,
            status='ESTORNADA',
            amount=10000,
            sale_date=date(2026, 7, 10),
            created_by=self.seller_user,
        )
        client = self._auth(self.manager, 'gestor123')
        seller = client.get(
            reverse('api-seller-dashboard-summary')
            + f'?period={self.period.uuid}',
        ).json()['sellers'][0]
        self.assertFalse(seller['has_sale_today'])
        self.assertFalse(seller['has_justification_today'])
        self.assertFalse(seller['has_day_resolved_today'])

    @patch('app.apps.api.views.timezone.localdate')
    def test_imported_sale_does_not_resolve_today(self, localdate_mock):
        localdate_mock.return_value = date(2026, 7, 10)
        Sale.objects.create(
            tenant=self.tenant,
            seller=self.seller,
            origin=Sale.Origin.IMPORTADA,
            status='ATIVA',
            amount=10000,
            sale_date=date(2026, 7, 10),
            created_by=self.manager,
        )
        client = self._auth(self.manager, 'gestor123')
        seller = client.get(
            reverse('api-seller-dashboard-summary')
            + f'?period={self.period.uuid}',
        ).json()['sellers'][0]
        self.assertFalse(seller['has_sale_today'])
        self.assertFalse(seller['has_justification_today'])
        self.assertFalse(seller['has_day_resolved_today'])

    @patch('app.apps.api.views.timezone.localdate')
    def test_other_tenant_justification_does_not_resolve_today(
        self, localdate_mock,
    ):
        localdate_mock.return_value = date(2026, 7, 10)
        other_user = User.objects.create_user(
            username='seller-other', role=User.Role.SELLER,
            tenant=self.tenant2,
        )
        other_seller = Seller.objects.create(
            tenant=self.tenant2, user=other_user, name='Outro vendedor',
            phone='11988887777', commission_rate=Decimal('0.01'),
        )
        SellerDayJustification.objects.create(
            tenant=self.tenant2,
            seller=other_seller,
            date=date(2026, 7, 10),
            reason=SellerDayJustification.Reason.FOLGA,
            created_by=self.manager2,
        )
        client = self._auth(self.manager, 'gestor123')
        sellers = client.get(
            reverse('api-seller-dashboard-summary')
            + f'?period={self.period.uuid}',
        ).json()['sellers']
        self.assertEqual([seller['uuid'] for seller in sellers], [
            str(self.seller.uuid),
        ])
        self.assertFalse(sellers[0]['has_justification_today'])
        self.assertFalse(sellers[0]['has_day_resolved_today'])

    def test_seller_dashboard_summary_queries_are_constant(self):
        client = self._auth(self.manager, 'gestor123')
        url = (
            reverse('api-seller-dashboard-summary')
            + f'?period={self.period.uuid}'
        )
        with CaptureQueriesContext(connection) as one_ctx:
            response = client.get(url)
        self.assertEqual(response.status_code, 200)

        for index in range(12):
            user = User.objects.create_user(
                username=f'bulk-user-{index}', role=User.Role.SELLER,
                tenant=self.tenant,
            )
            Seller.objects.create(
                tenant=self.tenant, user=user, name=f'Bulk {index}',
                phone=f'1198888{index:04d}',
                commission_rate=Decimal('0.01'),
            )

        with CaptureQueriesContext(connection) as thirteen_ctx:
            response = client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json()['sellers']), 13)

        for index in range(12, 49):
            user = User.objects.create_user(
                username=f'bulk-user-{index}', role=User.Role.SELLER,
                tenant=self.tenant,
            )
            Seller.objects.create(
                tenant=self.tenant, user=user, name=f'Bulk {index}',
                phone=f'1188888{index:04d}',
                commission_rate=Decimal('0.01'),
            )

        with CaptureQueriesContext(connection) as fifty_ctx:
            response = client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json()['sellers']), 50)
        self.assertEqual(
            len(thirteen_ctx.captured_queries),
            len(one_ctx.captured_queries),
        )
        self.assertEqual(
            len(fifty_ctx.captured_queries),
            len(one_ctx.captured_queries),
        )

    # helpers de selecao
    def test_resolve_selected_period_default_open(self):
        from app.apps.commissions.services import get_default_period
        # sem hoje coberto, cai na mais recente
        selected = get_default_period(self.tenant)
        self.assertEqual(selected.uuid, self.period.uuid)

    def test_get_previous_period(self):
        from app.apps.commissions.services import get_previous_period
        prev = CommissionPeriod.objects.create(
            tenant=self.tenant, month=6, year=2026,
            start_date=date(2026, 5, 21), end_date=date(2026, 6, 20),
        )
        result = get_previous_period(self.tenant, self.period)
        self.assertEqual(result.uuid, prev.uuid)

    # 21: sem competencia -> nenhum KPI financeiro por fallback silencioso
    def test_no_period_selector_empty(self):
        from app.apps.commissions.services import build_period_selector_context

        class _Req:
            GET = {}
        CommissionPeriod.objects.all().delete()
        ctx = build_period_selector_context(_Req(), self.tenant)
        self.assertFalse(ctx['has_periods'])
        self.assertIsNone(ctx['selected_period'])

    def test_two_periods_covering_today_raise_integrity_error(self):
        from app.apps.commissions.services import get_default_period, PeriodIntegrityError
        from django.utils import timezone

        today = timezone.localdate()
        second = CommissionPeriod.objects.create(
            tenant=self.tenant, month=1 if today.month == 12 else today.month + 1,
            year=today.year + 1 if today.month == 12 else today.year,
            start_date=today.replace(year=today.year + 2),
            end_date=today.replace(year=today.year + 2),
        )
        CommissionPeriod.objects.filter(uuid=second.uuid).update(
            start_date=today, end_date=today,
        )
        with self.assertRaises(PeriodIntegrityError):
            get_default_period(self.tenant)

    def test_other_tenant_period_in_html_view_returns_404(self):
        other_period = CommissionPeriod.objects.create(
            tenant=self.tenant2, month=8, year=2026,
            start_date=date(2026, 7, 21), end_date=date(2026, 8, 20),
        )
        self.client.force_login(self.manager)
        response = self.client.get(
            reverse('dashboard:gestor_vendedores') + f'?period={other_period.uuid}',
        )
        self.assertEqual(response.status_code, 404)

    def test_sale_change_log_count_is_one(self):
        sale = Sale.objects.get(sale_date=date(2026, 6, 25))
        SaleChangeLog.objects.create(
            sale=sale, tenant=self.tenant,
            action=SaleChangeLog.Action.UPDATE,
            changed_by=self.manager, reason='Correcao',
        )
        client = self._auth(self.manager, 'gestor123')
        response = client.get(
            f'/api/manager/seller/{self.seller.uuid}/'
            + f'?period={self.period.uuid}',
        )
        by_uuid = {item['uuid']: item for item in response.json()['manual_sales']}
        self.assertEqual(by_uuid[str(sale.uuid)]['change_log_count'], 1)


class StatementCompetenciaTest(TestCase):
    """LOTE 11 - extrato PDF por competencia + origem correta."""

    def setUp(self):
        cache.clear()
        self.tenant = Tenant.objects.create(
            company_name='Extrato Co', slug='extrato-co',
            default_commission_rate=Decimal('0.01'),
        )
        self.seller_user = User.objects.create_user(
            username='vendext', password='senha123',
            role=User.Role.SELLER, tenant=self.tenant,
        )
        self.seller = Seller.objects.create(
            tenant=self.tenant, name='Vend Extrato', phone='11999990000',
            user=self.seller_user, commission_rate=Decimal('0.01'),
        )
        self.period = CommissionPeriod.objects.create(
            tenant=self.tenant, month=7, year=2026,
            start_date=date(2026, 6, 21), end_date=date(2026, 7, 20),
            label='Julho/2026',
        )
        Sale.objects.create(
            tenant=self.tenant, seller=self.seller, origin=Sale.Origin.IMPORTADA,
            amount=100000, sale_date='2026-06-25', created_by=self.seller_user,
        )
        Sale.objects.create(
            tenant=self.tenant, seller=self.seller, origin=Sale.Origin.MANUAL,
            amount=200000, sale_date='2026-07-05', created_by=self.seller_user,
        )
        Sale.objects.create(
            tenant=self.tenant, seller=self.seller, origin=Sale.Origin.MANUAL,
            amount=999999, sale_date='2026-07-25', created_by=self.seller_user,
        )

    def _auth(self):
        client = APIClient()
        resp = client.post(reverse('api-login'), {
            'username': 'vendext', 'password': 'senha123',
        }, format='json')
        client.credentials(HTTP_AUTHORIZATION=f'Bearer {resp.data["access"]}')
        return client

    # 8: PDF total e comissao usam o mesmo range
    def test_statement_pdf_period_range(self):
        client = self._auth()
        url = reverse('api-seller-statement-period')
        resp = client.get(url + f'?period={self.period.uuid}')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/pdf')

    # 9: venda IMPORTADA aparece como Importada (nao Manual) - via context builder
    def test_importada_origin_display(self):
        s = Sale.objects.get(origin=Sale.Origin.IMPORTADA)
        self.assertEqual(s.get_origin_display(), 'Importada')


class ManagerCompetenciaExportTest(TestCase):
    """LOTE 12 - exports do gestor localizam SellerCommission pelo periodo."""

    def setUp(self):
        cache.clear()
        self.tenant = Tenant.objects.create(
            company_name='Export Co', slug='export-co',
            default_commission_rate=Decimal('0.01'),
        )
        self.manager = User.objects.create_user(
            username='gestorexp', password='gestor123',
            role=User.Role.MANAGER, tenant=self.tenant,
        )
        self.seller_user = User.objects.create_user(
            username='vendexp', password='senha123',
            role=User.Role.SELLER, tenant=self.tenant,
        )
        self.seller = Seller.objects.create(
            tenant=self.tenant, name='Vend Export', phone='11999997777',
            user=self.seller_user, commission_rate=Decimal('0.01'),
        )
        self.period = CommissionPeriod.objects.create(
            tenant=self.tenant, month=7, year=2026,
            start_date=date(2026, 6, 21), end_date=date(2026, 7, 20),
            label='Julho/2026', status=CommissionPeriod.Status.FECHADA,
        )
        self.sc = SellerCommission.objects.create(
            period=self.period, seller=self.seller,
            commission_rate=Decimal('0.01'),
            status=SellerCommission.Status.FECHADA,
            total_sold_amount=300000, commission_amount=3000,
        )
        Sale.objects.create(
            tenant=self.tenant, seller=self.seller, origin=Sale.Origin.MANUAL,
            amount=100000, sale_date='2026-06-25', created_by=self.seller_user,
        )
        Sale.objects.create(
            tenant=self.tenant, seller=self.seller, origin=Sale.Origin.MANUAL,
            amount=200000, sale_date='2026-07-05', created_by=self.seller_user,
        )
        Sale.objects.create(
            tenant=self.tenant, seller=self.seller, origin=Sale.Origin.IMPORTADA,
            amount=50000, sale_date='2026-07-10', created_by=self.seller_user,
        )
        Sale.objects.create(
            tenant=self.tenant, seller=self.seller, origin=Sale.Origin.IMPORTADA,
            amount=999999, sale_date='2026-07-21', created_by=self.seller_user,
        )

    def _auth(self):
        client = APIClient()
        resp = client.post(reverse('api-login'), {
            'username': 'gestorexp', 'password': 'gestor123',
        }, format='json')
        client.credentials(HTTP_AUTHORIZATION=f'Bearer {resp.data["access"]}')
        return client

    def test_csv_uses_frozen_commission_of_period(self):
        client = self._auth()
        url = f'/api/manager/seller/{self.seller.uuid}/csv/'
        resp = client.get(url + f'?period={self.period.uuid}')
        content = resp.content.decode('utf-8')
        # comissao congelada de 3000 centavos = 30.00
        self.assertIn('30.00', content)
        self.assertIn('2026-06-25', content)
        self.assertIn('2026-07-05', content)

    def test_csv_period_contains_only_period_sales(self):
        client = self._auth()
        response = client.get(
            reverse('api-seller-report-csv', args=[self.seller.uuid])
            + f'?period={self.period.uuid}',
        )
        content = response.content.decode('utf-8')
        self.assertIn('2026-06-25', content)
        self.assertIn('2026-07-05', content)
        self.assertIn('2026-07-10', content)
        self.assertNotIn('2026-07-21', content)

    def test_xlsx_period_contains_same_sales(self):
        from openpyxl import load_workbook

        client = self._auth()
        response = client.get(
            reverse('api-seller-report-xlsx', args=[self.seller.uuid])
            + f'?period={self.period.uuid}',
        )
        workbook = load_workbook(io.BytesIO(response.content), data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        rendered = '\n'.join(str(value) for row in rows for value in row if value)
        self.assertIn('25/06/2026', rendered)
        self.assertIn('05/07/2026', rendered)
        self.assertIn('10/07/2026', rendered)
        self.assertNotIn('21/07/2026', rendered)

    def test_pdf_period_contains_sales_label_range_and_imported_origin(self):
        client = self._auth()
        with patch('weasyprint.HTML') as html_class:
            html_class.return_value.write_pdf.return_value = b'%PDF-1.4'
            response = client.get(
                reverse('api-seller-report-pdf', args=[self.seller.uuid])
                + f'?period={self.period.uuid}',
            )
        self.assertEqual(response.status_code, 200)
        html = html_class.call_args.kwargs['string']
        self.assertIn(self.period.display_label, html)
        self.assertIn('21/06/2026', html)
        self.assertIn('20/07/2026', html)
        self.assertIn('25/06/2026', html)
        self.assertIn('05/07/2026', html)
        self.assertIn('10/07/2026', html)
        self.assertNotIn('21/07/2026', html)
        self.assertIn('Importada', html)

    def test_custom_range_exports_do_not_show_official_commission(self):
        client = self._auth()
        query = '?start=2026-06-25&end=2026-07-05'
        notice = 'Comissão oficial disponível apenas por competência'

        csv_response = client.get(
            reverse('api-seller-report-csv', args=[self.seller.uuid]) + query,
        )
        csv_content = csv_response.content.decode('utf-8')
        self.assertIn(notice, csv_content)
        self.assertNotIn('Comissao calculada', csv_content)

        from openpyxl import load_workbook
        xlsx_response = client.get(
            reverse('api-seller-report-xlsx', args=[self.seller.uuid]) + query,
        )
        workbook = load_workbook(io.BytesIO(xlsx_response.content), data_only=True)
        xlsx_text = '\n'.join(
            str(value)
            for row in workbook.active.iter_rows(values_only=True)
            for value in row if value
        )
        self.assertIn(notice, xlsx_text)
        self.assertNotIn('Comissao calculada', xlsx_text)

        with patch('weasyprint.HTML') as html_class:
            html_class.return_value.write_pdf.return_value = b'%PDF-1.4'
            pdf_response = client.get(
                reverse('api-seller-report-pdf', args=[self.seller.uuid]) + query,
            )
        self.assertEqual(pdf_response.status_code, 200)
        pdf_html = html_class.call_args.kwargs['string']
        self.assertIn(notice, pdf_html)
        self.assertNotIn('Comissao calculada', pdf_html)
