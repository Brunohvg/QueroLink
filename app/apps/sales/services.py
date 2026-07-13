import csv
import io
import os
import unicodedata
from datetime import datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from uuid import UUID

from django.core.cache import cache
from django.db import transaction
from django.utils import timezone
from app.apps.sales.models import Sale, SaleChangeLog
from app.apps.commissions.services import (
    resolve_period_for_date,
    ensure_seller_commission,
)
from app.apps.commissions.models import SellerCommission


EDITABLE_FIELDS = {'amount', 'sale_date', 'notes'}


def _validate_period(sale, ref_date):
    period = resolve_period_for_date(sale.tenant, ref_date)
    if not period:
        return True, None
    sc = SellerCommission.objects.filter(
        period=period, seller=sale.seller,
    ).first()
    if not sc:
        return True, None
    if not sc.is_editable:
        return False, (
            f'A comissão de {sale.seller.name} já foi fechada '
            f'ou paga nesta competência.'
        )
    return True, None


def update_sale_as_manager(sale, user, data, reason):
    if sale.status == 'ESTORNADA':
        raise ValueError('Venda estornada não pode ser alterada.')

    if not reason or len(reason.strip()) < 5:
        raise ValueError('Motivo é obrigatório e deve ter no mínimo 5 caracteres.')

    can, err = _validate_period(sale, sale.sale_date)
    if not can:
        raise ValueError(err)

    new_date = data.get('sale_date', sale.sale_date)
    if new_date != sale.sale_date:
        can, err = _validate_period(sale, new_date)
        if not can:
            raise ValueError(err)

    field_changes = {}
    for field in EDITABLE_FIELDS:
        if field in data:
            new_value = data[field]
            old_value = getattr(sale, field)
            if new_value != old_value:
                if isinstance(old_value, (int, float)):
                    field_changes[field] = {
                        'old': old_value,
                        'new': new_value,
                    }
                elif field == 'sale_date':
                    field_changes[field] = {
                        'old': str(old_value),
                        'new': str(new_value),
                    }
                else:
                    field_changes[field] = {
                        'old': str(old_value) if old_value else '',
                        'new': str(new_value) if new_value else '',
                    }

    if not field_changes:
        return sale, False

    with transaction.atomic():
        for field in EDITABLE_FIELDS:
            if field in data:
                setattr(sale, field, data[field])
        sale.updated_by = user
        sale.save()

        SaleChangeLog.objects.create(
            sale=sale,
            tenant=sale.tenant,
            action=SaleChangeLog.Action.UPDATE,
            changed_by=user,
            field_changes=field_changes,
            reason=reason.strip(),
        )

        if 'sale_date' in field_changes or 'amount' in field_changes:
            ensure_seller_commission(sale.seller, sale.sale_date)

    return sale, True


MAX_IMPORT_ROWS = 500
MAX_IMPORT_FILE_SIZE = 5 * 1024 * 1024
MAX_IMPORT_AMOUNT_CENTS = 10_000_000
IMPORT_PREVIEW_TTL = 30 * 60

CLASSIFICATION_SALE = 'SALE'
CLASSIFICATION_JUSTIFICATION = 'JUSTIFICATION_SUGGESTION'
CLASSIFICATION_EMPTY = 'EMPTY'
CLASSIFICATION_INVALID = 'INVALID'
CLASSIFICATION_DUPLICATE = 'DUPLICATE'
CLASSIFICATION_CONFLICT = 'CONFLICT'
CLASSIFICATION_BLOCKED = 'BLOCKED'

KNOWN_REASONS = {
    'ATESTADO': 'ATESTADO',
    'FALTA': 'FALTA',
    'FERIAS': 'FERIAS',
    'FOLGA': 'FOLGA',
    'AFASTAMENTO': 'AFASTAMENTO',
    'FERIADO': 'FERIADO',
    'SEM EXPEDIENTE': 'SEM_EXPEDIENTE',
    'OUTRO': 'OUTRO',
}


def _preview_cache_key(batch_uuid):
    return f'sales-import-preview:{batch_uuid}'


def _normalize_text(value):
    text = unicodedata.normalize('NFKD', str(value or '').strip())
    text = ''.join(char for char in text if not unicodedata.combining(char))
    return ' '.join(text.replace('_', ' ').upper().split())


def _parse_amount_cents(value):
    raw = str(value or '').strip()
    if not raw:
        return None
    normalized = raw.replace('R$', '').replace('\u00a0', '').replace(' ', '')
    if normalized.startswith('-') or normalized.startswith('+'):
        return None
    if not normalized or any(
        char not in '0123456789.,' for char in normalized
    ):
        return None

    if ',' in normalized:
        if normalized.count(',') != 1:
            return None
        integer, decimal = normalized.rsplit(',', 1)
        if len(decimal) > 2 or not decimal:
            return None
        integer = integer.replace('.', '')
        canonical = f'{integer}.{decimal}'
    elif '.' in normalized:
        if normalized.count('.') == 1:
            integer, decimal = normalized.split('.', 1)
            if not decimal or len(decimal) > 2:
                return None
            canonical = normalized
        else:
            return None
    else:
        canonical = normalized

    try:
        amount = Decimal(canonical)
    except InvalidOperation:
        return None
    cents = amount * 100
    if cents != cents.to_integral_value():
        return None
    cents = int(cents)
    if cents <= 0 or cents > MAX_IMPORT_AMOUNT_CENTS:
        return None
    return cents


def _parse_br_number(value):
    """Compatibilidade interna: retorna Decimal sem usar float."""
    cents = _parse_amount_cents(value)
    return Decimal(cents) / 100 if cents is not None else None


def _parse_date(value):
    if not value:
        return None
    s = str(value).strip()
    for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _parse_csv(content):
    delim = ','
    first_line = content.split('\n')[0] if '\n' in content else content
    if first_line.count(';') > first_line.count(','):
        delim = ';'
    try:
        reader = csv.reader(io.StringIO(content), delimiter=delim, strict=True)
        rows = []
        for row in reader:
            if row:
                rows.append(row)
        return rows
    except csv.Error as exc:
        raise ValueError('Arquivo CSV invalido ou corrompido.') from exc


def _parse_xlsx(content_bytes):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(
            io.BytesIO(content_bytes), read_only=True, data_only=True,
        )
        worksheets = [sheet for sheet in wb.worksheets if sheet.max_row]
        if not worksheets:
            return []
        ws = worksheets[0]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([str(c) if c is not None else '' for c in row])
        return rows
    except ImportError:
        raise ValueError('Suporte XLSX indisponivel.')
    except Exception as exc:
        raise ValueError('Arquivo XLSX invalido ou corrompido.') from exc


def normalize_headers(headers):
    mapping = {}
    for i, h in enumerate(headers):
        key = _normalize_text(h).lower()
        if key in ('data', 'date', 'dt'):
            mapping['date'] = i
        elif key in ('vendedor', 'vendedora', 'seller', 'nome'):
            mapping['seller'] = i
        elif key in ('valor', 'value', 'total', 'amount', 'ocorrencia'):
            mapping['amount'] = i
        elif key in ('observacao', 'obs', 'observacoes', 'notes'):
            mapping['notes'] = i
        elif key in ('seller uuid', 'vendedor uuid', 'uuid vendedor', 'uuid'):
            mapping['seller_uuid'] = i
        elif key in ('codigo_vendedor', 'codigo', 'code', 'id'):
            mapping['codigo_vendedor'] = i
    return mapping


def _row_value(row, index):
    return row[index] if index is not None and index < len(row) else ''


def _legacy_status(classification, requires_manual_action=False):
    if classification == CLASSIFICATION_SALE:
        return 'ok'
    if classification == CLASSIFICATION_DUPLICATE:
        return 'duplicate'
    if classification == CLASSIFICATION_CONFLICT and requires_manual_action:
        return 'needs_selection'
    return 'error'


def classify_import_rows(tenant, rows, header_map, seller_choices=None):
    """Classifica todas as linhas usando consultas em lote e escopo do tenant."""
    from app.apps.commissions.models import CommissionPeriod, SellerCommission
    from app.apps.sellers.models import Seller, SellerDayJustification

    seller_choices = {
        str(key): str(value) for key, value in (seller_choices or {}).items()
    }
    sellers = list(
        Seller.objects.filter(tenant=tenant).select_related('user')
    )
    sellers_by_uuid = {str(s.uuid): s for s in sellers}
    sellers_by_username = {}
    sellers_by_name = {}
    for seller in sellers:
        if seller.user_id:
            sellers_by_username.setdefault(seller.user.username.casefold(), []).append(seller)
        sellers_by_name.setdefault(seller.name.strip().casefold(), []).append(seller)

    prepared = []
    valid_dates = []
    for idx, row in enumerate(rows, start=1):
        date_raw = _row_value(row, header_map.get('date'))
        value_raw = _row_value(row, header_map.get('amount'))
        seller_raw = _row_value(row, header_map.get('seller'))
        uuid_raw = _row_value(row, header_map.get('seller_uuid'))
        code_raw = _row_value(row, header_map.get('codigo_vendedor'))
        notes = str(_row_value(row, header_map.get('notes')) or '')[:255]
        parsed_date = _parse_date(date_raw)
        if parsed_date:
            valid_dates.append(parsed_date)
        prepared.append({
            'row_number': idx,
            'row': row,
            'date_raw': date_raw,
            'raw_value': str(value_raw or '').strip(),
            'seller_raw': str(seller_raw or '').strip(),
            'uuid_raw': str(uuid_raw or '').strip(),
            'code_raw': str(code_raw or '').strip(),
            'notes': notes,
            'date': parsed_date,
        })

    periods = []
    sales = []
    justifications = []
    commissions = []
    if valid_dates:
        start_date, end_date = min(valid_dates), max(valid_dates)
        periods = list(CommissionPeriod.objects.filter(
            tenant=tenant,
            start_date__lte=end_date,
            end_date__gte=start_date,
        ))
        sales = list(Sale.objects.filter(
            tenant=tenant,
            sale_date__range=(start_date, end_date),
            status='ATIVA',
        ))
        justifications = list(SellerDayJustification.objects.filter(
            tenant=tenant, date__range=(start_date, end_date),
        ))
        commissions = list(SellerCommission.objects.filter(
            period__in=periods, seller__tenant=tenant,
        ))

    sales_by_day = {}
    for sale in sales:
        sales_by_day.setdefault((sale.seller_id, sale.sale_date), []).append(sale)
    justifications_by_day = {
        (item.seller_id, item.date): item for item in justifications
    }
    commission_by_pair = {
        (item.period_id, item.seller_id): item for item in commissions
    }

    results = []
    for item in prepared:
        classification = CLASSIFICATION_INVALID
        message = ''
        suggestions = []
        requires_manual = False
        seller = None
        match_type = None
        period = None
        existing_sale = None
        existing_justification = None
        normalized_amount = _parse_amount_cents(item['raw_value'])
        normalized_reason = KNOWN_REASONS.get(_normalize_text(item['raw_value']))

        if not item['raw_value']:
            classification = CLASSIFICATION_EMPTY
            message = 'Linha sem valor ou ocorrencia.'
        elif not item['date']:
            message = f'Data invalida: {item["date_raw"]}'
        else:
            choice = seller_choices.get(str(item['row_number']))
            if choice:
                seller = sellers_by_uuid.get(choice)
                match_type = 'manual'
                if not seller:
                    message = 'Vendedor selecionado nao pertence ao tenant.'
            if not seller and not choice and item['uuid_raw']:
                try:
                    seller = sellers_by_uuid.get(str(UUID(item['uuid_raw'])))
                except ValueError:
                    seller = None
                match_type = 'uuid' if seller else None
            identifier = item['code_raw'] or item['seller_raw']
            if not seller and not choice and identifier:
                username_matches = sellers_by_username.get(identifier.casefold(), [])
                name_matches = sellers_by_name.get(identifier.casefold(), [])
                matches = username_matches or name_matches
                if len(matches) == 1:
                    seller = matches[0]
                    match_type = 'username' if username_matches else 'name_exact'
                elif len(matches) > 1:
                    suggestions = [
                        {'uuid': str(candidate.uuid), 'name': candidate.name}
                        for candidate in matches
                    ]
                    classification = CLASSIFICATION_CONFLICT
                    requires_manual = True
                    message = 'Mais de um vendedor corresponde ao nome informado.'
            if not seller and not message:
                partial = [
                    candidate for candidate in sellers
                    if identifier.casefold() in candidate.name.casefold()
                ][:5] if identifier else []
                suggestions = [
                    {'uuid': str(candidate.uuid), 'name': candidate.name}
                    for candidate in partial
                ]
                if suggestions:
                    classification = CLASSIFICATION_CONFLICT
                    requires_manual = True
                    message = 'Selecione manualmente um vendedor.'
                else:
                    message = (
                        f'Vendedor nao encontrado: {identifier}'
                        if identifier else 'Vendedor nao informado.'
                    )

            if seller:
                covering = [
                    candidate for candidate in periods
                    if candidate.start_date <= item['date'] <= candidate.end_date
                ]
                if len(covering) > 1:
                    classification = CLASSIFICATION_CONFLICT
                    message = 'Mais de uma competencia cobre esta data.'
                elif not covering:
                    classification = CLASSIFICATION_BLOCKED
                    message = 'Nenhuma competencia cobre esta data.'
                else:
                    period = covering[0]
                    locked_statuses = {
                        CommissionPeriod.Status.FECHADA,
                        CommissionPeriod.Status.PAGA,
                        CommissionPeriod.Status.CANCELADA,
                    }
                    seller_commission = commission_by_pair.get(
                        (period.pk, seller.pk),
                    )
                    if (
                        period.status in locked_statuses
                        or (seller_commission and not seller_commission.is_editable)
                    ):
                        classification = CLASSIFICATION_BLOCKED
                        message = f'A competencia {period.display_label} esta bloqueada.'
                    else:
                        day_sales = sales_by_day.get((seller.pk, item['date']), [])
                        existing_justification = justifications_by_day.get(
                            (seller.pk, item['date']),
                        )
                        if normalized_amount is not None:
                            existing_sale = next(
                                (sale for sale in day_sales if sale.amount == normalized_amount),
                                None,
                            )
                            if existing_justification:
                                classification = CLASSIFICATION_CONFLICT
                                message = (
                                    'Este dia ja possui justificativa: '
                                    f'{existing_justification.get_reason_display()}.'
                                )
                            elif existing_sale:
                                classification = CLASSIFICATION_DUPLICATE
                                message = 'Venda duplicada (mesmo vendedor, data e valor).'
                            else:
                                classification = CLASSIFICATION_SALE
                        elif normalized_reason:
                            active_sales = [
                                sale for sale in day_sales
                                if sale.origin in Sale.COMMISSION_ORIGINS
                            ]
                            existing_sale = active_sales[0] if active_sales else None
                            if existing_sale:
                                classification = CLASSIFICATION_CONFLICT
                                message = 'Este dia ja possui venda registrada.'
                            elif existing_justification:
                                if existing_justification.reason == normalized_reason:
                                    classification = CLASSIFICATION_DUPLICATE
                                    message = 'Justificativa identica ja cadastrada.'
                                else:
                                    classification = CLASSIFICATION_CONFLICT
                                    message = (
                                        'Este dia ja possui justificativa: '
                                        f'{existing_justification.get_reason_display()}.'
                                    )
                            else:
                                classification = CLASSIFICATION_JUSTIFICATION
                                message = 'Justificativa sugerida; exige confirmacao.'
                        else:
                            classification = CLASSIFICATION_INVALID
                            message = f'Valor ou ocorrencia invalida: {item["raw_value"]}'

        result = {
            'row_number': item['row_number'],
            'seller_identifier_original': (
                item['uuid_raw'] or item['code_raw'] or item['seller_raw']
            ),
            'seller_uuid': str(seller.uuid) if seller else None,
            'seller_name': seller.name if seller else None,
            'date': item['date'].isoformat() if item['date'] else None,
            'raw_value': item['raw_value'],
            'normalized_amount': normalized_amount,
            'normalized_reason': normalized_reason,
            'classification': classification,
            'message': message,
            'period_uuid': str(period.uuid) if period else None,
            'period_label': period.display_label if period else None,
            'existing_sale_uuid': str(existing_sale.uuid) if existing_sale else None,
            'existing_justification_uuid': (
                str(existing_justification.uuid) if existing_justification else None
            ),
            'requires_manual_action': requires_manual,
            'suggestions': suggestions,
            'match_type': match_type,
            'notes': item['notes'],
        }
        result.update({
            'status': _legacy_status(classification, requires_manual),
            'resolved_seller': result['seller_uuid'],
            'resolved_seller_uuid': result['seller_uuid'],
            'resolved_seller_name': result['seller_name'],
            'sale_date': result['date'],
            'amount_cents': result['normalized_amount'],
        })
        results.append(result)
    return results


def preview_import_rows(tenant, rows, header_map):
    return classify_import_rows(tenant, rows, header_map)


def create_import_preview(*, tenant, user, file_obj):
    from app.apps.sales.models import SaleImportBatch

    if not file_obj:
        raise ValueError('Arquivo obrigatorio.')
    if file_obj.size > MAX_IMPORT_FILE_SIZE:
        raise ValueError('Arquivo muito grande. Maximo 5MB.')
    filename = os.path.basename(str(file_obj.name or ''))[:255]
    extension = os.path.splitext(filename)[1].lower()
    if extension not in ('.csv', '.xlsx'):
        raise ValueError('Formato nao suportado. Envie CSV ou XLSX.')
    content_bytes = file_obj.read()
    if not content_bytes:
        raise ValueError('Arquivo vazio ou sem dados.')
    file_hash = sha256(content_bytes).hexdigest()

    if extension == '.csv':
        try:
            rows = _parse_csv(content_bytes.decode('utf-8-sig'))
        except UnicodeDecodeError as exc:
            raise ValueError('CSV deve estar em UTF-8, com ou sem BOM.') from exc
    else:
        rows = _parse_xlsx(content_bytes)
    if not rows or len(rows) < 2:
        raise ValueError('Arquivo vazio ou sem dados.')
    headers, data_rows = rows[0], rows[1:]
    if len(data_rows) > MAX_IMPORT_ROWS:
        raise ValueError('Maximo 500 linhas por importacao.')
    header_map = normalize_headers(headers)
    missing = [
        label for key, label in (
            ('date', 'data'), ('seller', 'vendedor'), ('amount', 'valor ou ocorrencia'),
        ) if key not in header_map
    ]
    if missing:
        raise ValueError(
            f'Cabecalhos obrigatorios nao encontrados: {", ".join(missing)}.'
        )

    already_imported = SaleImportBatch.objects.filter(
        tenant=tenant, file_hash=file_hash, status='IMPORTED',
    ).exists()
    results = classify_import_rows(tenant, data_rows, header_map)
    batch = SaleImportBatch.objects.create(
        tenant=tenant,
        filename=filename,
        file_hash=file_hash,
        uploaded_by=user,
        status='PENDING',
        total_rows=len(data_rows),
        duplicate_count=sum(
            item['classification'] == CLASSIFICATION_DUPLICATE for item in results
        ),
        error_count=sum(
            item['classification'] in (
                CLASSIFICATION_INVALID, CLASSIFICATION_CONFLICT, CLASSIFICATION_BLOCKED,
            ) for item in results
        ),
    )
    cache.set(_preview_cache_key(batch.uuid), {
        'rows': data_rows,
        'header_map': header_map,
        'file_hash': file_hash,
    }, IMPORT_PREVIEW_TTL)
    counts = {
        key: sum(item['classification'] == key for item in results)
        for key in (
            CLASSIFICATION_SALE, CLASSIFICATION_JUSTIFICATION,
            CLASSIFICATION_EMPTY, CLASSIFICATION_INVALID,
            CLASSIFICATION_DUPLICATE, CLASSIFICATION_CONFLICT,
            CLASSIFICATION_BLOCKED,
        )
    }
    return {
        'batch_uuid': str(batch.uuid),
        'filename': filename,
        'file_hash': file_hash,
        'total_rows': len(data_rows),
        'headers': headers,
        'results': results,
        'already_imported': already_imported,
        'counts': counts,
        'ok_count': counts[CLASSIFICATION_SALE],
        'justification_count': counts[CLASSIFICATION_JUSTIFICATION],
        'duplicate_count': counts[CLASSIFICATION_DUPLICATE],
        'conflict_count': counts[CLASSIFICATION_CONFLICT],
        'blocked_count': counts[CLASSIFICATION_BLOCKED],
        'empty_count': counts[CLASSIFICATION_EMPTY],
        'error_count': counts[CLASSIFICATION_INVALID],
        'needs_selection_count': sum(
            item['requires_manual_action'] for item in results
        ),
        'total_amount': sum(
            item['normalized_amount'] or 0 for item in results
            if item['classification'] == CLASSIFICATION_SALE
        ),
        'expires_in_seconds': IMPORT_PREVIEW_TTL,
    }


def import_sales(tenant, user, batch_uuid, decisions=None, force_reimport=False):
    """Confirma um preview preservado; nunca recebe dados financeiros do cliente."""
    from app.apps.sales.models import SaleImportBatch
    from app.apps.sellers.models import Seller
    from app.apps.sellers.services import create_day_justification

    decisions = decisions or {}
    seller_choices = decisions.get('seller_choices') or {}
    confirmed_justifications = {
        int(value) for value in decisions.get('confirmed_justifications', [])
        if str(value).isdigit()
    }
    justification_notes = decisions.get('justification_notes') or {}
    cached = cache.get(_preview_cache_key(batch_uuid))
    if not cached:
        raise ValueError('Preview expirado ou inexistente. Envie o arquivo novamente.')

    created = 0
    justifications_created = 0
    duplicates = 0
    conflicts = 0
    total_amount = 0
    errors = []
    with transaction.atomic():
        batch = SaleImportBatch.objects.select_for_update().get(
            uuid=batch_uuid, tenant=tenant, uploaded_by=user, status='PENDING',
        )
        if cached.get('file_hash') != batch.file_hash:
            raise ValueError('Hash do preview invalido.')
        if SaleImportBatch.objects.filter(
            tenant=tenant, file_hash=batch.file_hash, status='IMPORTED',
        ).exclude(pk=batch.pk).exists() and not force_reimport:
            raise ValueError('Este arquivo ja foi importado. Confirme a reimportacao.')

        results = classify_import_rows(
            tenant, cached['rows'], cached['header_map'], seller_choices,
        )
        sellers = {
            str(seller.uuid): seller
            for seller in Seller.objects.filter(
                tenant=tenant,
                uuid__in=[item['seller_uuid'] for item in results if item['seller_uuid']],
            )
        }
        for item in results:
            classification = item['classification']
            row_number = item['row_number']
            if classification == CLASSIFICATION_DUPLICATE:
                duplicates += 1
                continue
            if classification in (CLASSIFICATION_CONFLICT, CLASSIFICATION_BLOCKED):
                conflicts += 1
                continue
            if classification not in (CLASSIFICATION_SALE, CLASSIFICATION_JUSTIFICATION):
                continue
            if (
                classification == CLASSIFICATION_JUSTIFICATION
                and row_number not in confirmed_justifications
            ):
                continue
            try:
                seller = sellers[item['seller_uuid']]
                sale_date = _parse_date(item['date'])
                with transaction.atomic():
                    if classification == CLASSIFICATION_SALE:
                        sale = Sale.objects.create(
                            tenant=tenant,
                            seller=seller,
                            origin=Sale.Origin.IMPORTADA,
                            amount=item['normalized_amount'],
                            sale_date=sale_date,
                            notes=item['notes'],
                            created_by=user,
                        )
                        SaleChangeLog.objects.create(
                            sale=sale,
                            tenant=tenant,
                            action=SaleChangeLog.Action.CREATE_IMPORT,
                            changed_by=user,
                            field_changes={
                                'import_batch': str(batch.uuid),
                                'filename': batch.filename,
                                'amount': item['normalized_amount'],
                                'sale_date': item['date'],
                            },
                            reason=f'Importado via planilha: {batch.filename}',
                        )
                        ensure_seller_commission(seller, sale_date)
                        created += 1
                        total_amount += item['normalized_amount']
                    else:
                        notes = str(
                            justification_notes.get(str(row_number), '') or ''
                        )[:255]
                        create_day_justification(
                            tenant=tenant,
                            seller=seller,
                            date=sale_date,
                            reason=item['normalized_reason'],
                            notes=notes,
                            user=user,
                        )
                        justifications_created += 1
            except Exception as exc:
                errors.append({'row_number': row_number, 'error': str(exc)})

        batch.status = 'IMPORTED' if created or justifications_created else 'REJECTED'
        batch.created_count = created
        batch.duplicate_count = duplicates
        batch.error_count = len(errors) + conflicts
        batch.save(update_fields=[
            'status', 'created_count', 'duplicate_count', 'error_count',
        ])

    cache.delete(_preview_cache_key(batch_uuid))
    return {
        'batch_uuid': str(batch.uuid),
        'file_hash': batch.file_hash,
        'filename': batch.filename,
        'created': created,
        'justifications_created': justifications_created,
        'duplicates': duplicates,
        'conflicts': conflicts,
        'errors': errors,
        'total': batch.total_rows,
        'total_amount': total_amount,
        'user': user.get_username(),
        'confirmed_at': timezone.now().isoformat(),
        'results': results,
    }
