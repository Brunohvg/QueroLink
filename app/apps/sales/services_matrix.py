import io
import uuid
import logging
from datetime import date, timedelta
from calendar import monthrange

from django.conf import settings
from django.utils import timezone

from app.apps.sellers.models import Seller as SellerModel

logger = logging.getLogger(__name__)

KNOWN_SITUATIONS = {'FERIAS', 'FALTA', 'ATESTADO', 'FOLGA', 'SEM EXPEDIENTE'}

MONTHS_PT = [
    '', 'Janeiro', 'Fevereiro', 'Marco', 'Abril', 'Maio', 'Junho',
    'Julho', 'Agosto', 'Setembro', 'Outubro', 'Novembro', 'Dezembro',
]


def _str_to_float_br(val):
    """Converte valor brasileiro (12.917,25 ou 12917,25) para float."""
    if isinstance(val, (int, float)):
        return float(val)
    val = str(val).strip().replace('R$', '').replace(' ', '').strip()
    if '.' in val and ',' in val:
        if val.rindex(',') > val.rindex('.'):
            val = val.replace('.', '')  # 12.917,25 → 12917,25
            val = val.replace(',', '.')
        else:
            val = val.replace(',', '')  # 12,917.25 → 12917.25
    elif ',' in val:
        val = val.replace(',', '.')
    try:
        return round(float(val), 2)
    except (ValueError, TypeError):
        return None


def _normalize(text):
    import unicodedata
    return unicodedata.normalize('NFKD', str(text or '')).encode('ascii', errors='ignore').decode().strip().upper()


def generate_template_xlsx(tenant, start_date, end_date):
    """Gera planilha XLSX com grade de vendas."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, numbers
    from openpyxl.utils import get_column_letter

    sellers = list(SellerModel.objects.filter(
        tenant=tenant, is_active=True,
    ).order_by('name'))

    dates = []
    d = start_date
    while d <= end_date:
        dates.append(d)
        d += timedelta(days=1)

    wb = Workbook()

    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color='E5E7EB', end_color='E5E7EB', fill_type='solid')
    weekend_fill = PatternFill(start_color='F3F4F6', end_color='F3F4F6', fill_type='solid')

    # ── Aba IMPORTACAO ──────────────────────────────────────────
    ws = wb.active
    ws.title = 'IMPORTACAO'

    # Header row
    ws.cell(row=1, column=1, value='DATA').font = header_font
    ws.cell(row=1, column=1).fill = header_fill
    ws.cell(row=1, column=1).border = thin_border
    seller_col_map = {}
    for i, s in enumerate(sellers, start=2):
        cell = ws.cell(row=1, column=i, value=s.name)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='right')
        seller_col_map[i] = s

    # Data rows
    for row_idx, d in enumerate(dates, start=2):
        cell = ws.cell(row=row_idx, column=1, value=d)
        cell.number_format = 'DD/MM/AAAA'
        cell.border = thin_border
        if d.weekday() >= 5:
            cell.fill = weekend_fill

        for col_idx in range(2, 2 + len(sellers)):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.number_format = '#,##0.00'
            cell.border = thin_border
            if d.weekday() >= 5:
                cell.fill = weekend_fill

    # Freeze panes (header + date column)
    ws.freeze_panes = 'B2'

    # Column widths
    ws.column_dimensions['A'].width = 14
    for i in range(2, 2 + len(sellers)):
        ws.column_dimensions[get_column_letter(i)].width = 18

    # Auto-filter
    if dates:
        last_col = get_column_letter(1 + len(sellers))
        last_row = 1 + len(dates)
        ws.auto_filter.ref = f'A1:{last_col}{last_row}'

    # ── Aba OBSERVACOES ─────────────────────────────────────────
    ws2 = wb.create_sheet('OBSERVACOES')
    obs_headers = ['DATA', 'VENDEDOR', 'OBSERVACAO']
    for i, h in enumerate(obs_headers, start=1):
        cell = ws2.cell(row=1, column=i, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
    ws2.column_dimensions['A'].width = 14
    ws2.column_dimensions['B'].width = 22
    ws2.column_dimensions['C'].width = 60
    ws2.freeze_panes = 'A2'

    # ── Aba INSTRUCOES ──────────────────────────────────────────
    ws3 = wb.create_sheet('INSTRUCOES')
    instructions = [
        'INSTRUCOES PARA PREENCHIMENTO',
        '',
        f'Periodo: {start_date.strftime("%d/%m/%Y")} a {end_date.strftime("%d/%m/%Y")}',
        f'Gerado em: {timezone.now().strftime("%d/%m/%Y %H:%M")}',
        f'Tenant: {tenant.company_name}',
        '',
        '--- ABA IMPORTACAO ---',
        'Preencha os valores de venda de cada vendedor por data.',
        '',
        'Datas: use formato DD/MM/AAAA.',
        'Valores: use virgula como separador decimal (ex: 12917,25 ou 12.917,25).',
        'Situacoes permitidas: FERIAS, FALTA, ATESTADO, FOLGA, SEM EXPEDIENTE.',
        'Celula vazia: nao cria venda (ignorada).',
        '',
        '--- ABA OBSERVACOES ---',
        'Use para observacoes livres vinculadas a uma venda ou situacao.',
        'Nao use formulas, HTML ou macros.',
        '',
        '--- ANTES DE IMPORTAR ---',
        '1. Preencha a aba IMPORTACAO com valores ou situacoes.',
        '2. Preencha a aba OBSERVACOES se necessario.',
        '3. Faca upload do arquivo no sistema.',
        '4. Revise a previa antes de confirmar.',
        '5. A confirmacao e definitiva — use o recurso de desfazer se necessario.',
        '',
        'Versao do modelo: 1.0',
    ]
    for i, line in enumerate(instructions, start=1):
        cell = ws3.cell(row=i, column=1, value=line)
        if line and not line.startswith('---') and not line.startswith('Versao'):
            cell.font = Font(size=11)
        elif line.startswith('---'):
            cell.font = Font(bold=True, size=12)
        elif line == 'INSTRUCOES PARA PREENCHIMENTO':
            cell.font = Font(bold=True, size=14)
    ws3.column_dimensions['A'].width = 90

    # ── Aba _META (oculta) ──────────────────────────────────────
    ws4 = wb.create_sheet('_META')
    ws4.sheet_state = 'hidden'
    meta = [
        ('versao', '1.0'),
        ('tenant_uuid', str(tenant.uuid)),
        ('tenant_nome', tenant.company_name),
        ('periodo_inicio', start_date.isoformat()),
        ('periodo_fim', end_date.isoformat()),
        ('gerado_em', timezone.now().isoformat()),
    ]
    for i, s in enumerate(sellers, start=1):
        meta.append((f'vendedor_uuid_{i}', str(s.uuid)))
        meta.append((f'vendedor_nome_{i}', s.name))
    for i, (k, v) in enumerate(meta, start=1):
        ws4.cell(row=i, column=1, value=k)
        ws4.cell(row=i, column=2, value=v)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def parse_matrix_xlsx(content_bytes, tenant):
    """Parseia formato grade. Retorna lista de dicts normalizados."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content_bytes), read_only=True, data_only=True)
    if 'IMPORTACAO' not in wb.sheetnames:
        raise ValueError('Planilha deve conter aba IMPORTACAO.')

    ws = wb['IMPORTACAO']
    rows_iter = ws.iter_rows(values_only=False)

    # Ler cabecalhos (primeira linha)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise ValueError('Planilha vazia.')

    headers = [str(c.value).strip().upper() if c.value is not None else '' for c in header_row]
    if not headers or headers[0] != 'DATA':
        raise ValueError('Primeira coluna deve ser DATA.')

    seller_names = headers[1:]
    if not seller_names:
        raise ValueError('Nenhum vendedor encontrado no cabecalho.')

    # Validar colunas duplicadas
    seen = set()
    for name in seller_names:
        norm = _normalize(name)
        if norm in seen:
            raise ValueError(f'Coluna duplicada: {name}.')
        seen.add(norm)

    # Carregar metadados (se _META existir)
    meta_seller_uuids = {}
    meta_seller_names = {}
    if '_META' in wb.sheetnames:
        ws_meta = wb['_META']
        for row in ws_meta.iter_rows(values_only=True):
            if row[0] and row[1]:
                key = str(row[0]).strip()
                val = str(row[1]).strip()
                if key.startswith('vendedor_uuid_'):
                    idx = key.replace('vendedor_uuid_', '')
                    meta_seller_uuids[idx] = val
                elif key.startswith('vendedor_nome_'):
                    idx = key.replace('vendedor_nome_', '')
                    meta_seller_names[idx] = val

    # Buscar vendedores do tenant
    from app.apps.sellers.models import Seller
    active_sellers = {str(s.uuid): s for s in Seller.objects.filter(tenant=tenant, is_active=True)}
    all_tenant_sellers = {str(s.uuid): s for s in Seller.objects.filter(tenant=tenant)}

    # Mapear colunas para vendedores
    col_seller_map = [None] * len(seller_names)
    col_inactive = [False] * len(seller_names)

    # Primeiro: tentar mapear por UUID da _META
    for idx, name in enumerate(seller_names):
        meta_idx = str(idx + 1)
        if meta_idx in meta_seller_uuids:
            suuid = meta_seller_uuids[meta_idx]
            if suuid in all_tenant_sellers:
                col_seller_map[idx] = all_tenant_sellers[suuid]
                if suuid not in active_sellers:
                    col_inactive[idx] = True
                continue

    # Segundo: mapear por nome (exato, normalizado)
    name_map = {}
    for suuid, s in all_tenant_sellers.items():
        name_map[_normalize(s.name)] = s

    for idx, name in enumerate(seller_names):
        if col_seller_map[idx] is not None:
            continue
        norm = _normalize(name)
        if norm in name_map:
            s = name_map[norm]
            col_seller_map[idx] = s
            if str(s.uuid) not in active_sellers:
                col_inactive[idx] = True
        else:
            raise ValueError(f'Vendedor nao encontrado: {name}')

    # Limitar vendedores
    MAX_SELLERS = 100
    if len(col_seller_map) > MAX_SELLERS:
        raise ValueError(f'Maximo de {MAX_SELLERS} vendedores por planilha.')

    # Processar linhas
    results = []
    MAX_ROWS = 2000
    row_count = 0
    for row in rows_iter:
        row_count += 1
        if row_count > MAX_ROWS:
            break
        cells = [c.value for c in row]
        if not cells or cells[0] is None:
            continue  # linha vazia

        raw_date = cells[0]
        parsed_date = _parse_date_cell(raw_date)
        if parsed_date is None:
            continue  # data invalida

        for col_idx, seller in enumerate(col_seller_map):
            if seller is None:
                continue
            raw_val = cells[col_idx + 1] if col_idx + 1 < len(cells) else None
            if raw_val is None or (isinstance(raw_val, str) and raw_val.strip() == ''):
                continue

            result = {
                'date': parsed_date,
                'seller': seller,
                'seller_uuid': str(seller.uuid),
                'seller_name': seller.name,
                'inactive': col_inactive[col_idx],
            }

            str_val = str(raw_val).strip().upper()
            if str_val in KNOWN_SITUATIONS:
                result['type'] = 'justification'
                result['justification'] = str_val
                result['amount_cents'] = 0
            else:
                amount_float = _str_to_float_br(raw_val)
                if amount_float is None or amount_float <= 0:
                    raise ValueError(
                        f'Valor invalido na linha {row_count + 1}, coluna {seller.name}: '
                        f'"{raw_val}". Use numero, FERIAS, FALTA, ATESTADO, FOLGA ou SEM EXPEDIENTE.'
                    )
                from app.apps.sales.services import MAX_IMPORT_AMOUNT_CENTS
                amount_cents = int(round(amount_float * 100))
                if amount_cents > MAX_IMPORT_AMOUNT_CENTS:
                    raise ValueError(
                        f'Valor acima do limite em {seller.name} em {parsed_date}: R$ {amount_float:,.2f}'
                    )
                result['type'] = 'sale'
                result['amount_cents'] = amount_cents
                result['amount_float'] = amount_float

            results.append(result)

    wb.close()
    return results


def _parse_date_cell(val):
    """Converte celula de data para date."""
    if isinstance(val, date):
        return val
    if isinstance(val, (int, float)):
        from datetime import datetime
        try:
            return (
                datetime(1899, 12, 30) + timedelta(days=int(val))
            ).date()
        except (ValueError, OverflowError):
            return None
    s = str(val).strip()
    formats = ['%d/%m/%Y', '%Y-%m-%d']
    for fmt in formats:
        try:
            from datetime import datetime
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def detect_import_format(content_bytes, filename):
    """Detecta se arquivo e no formato grade ou lista. Retorna 'matrix', 'list' ou None."""
    if filename.lower().endswith('.csv'):
        return 'list'

    from openpyxl import load_workbook
    try:
        wb = load_workbook(
            io.BytesIO(content_bytes), read_only=True, data_only=True,
        )
    except Exception:
        return None

    sheet_names = wb.sheetnames
    if 'IMPORTACAO' in sheet_names:
        wb.close()
        return 'matrix'

    # Sem aba IMPORTACAO, tenta detectar pelo cabecalho na aba ativa
    try:
        ws = wb.active
        header_row = next(ws.iter_rows(values_only=True))
    except StopIteration:
        wb.close()
        return None

    if not header_row or not header_row[0]:
        wb.close()
        return None

    first_col = str(header_row[0]).strip().upper()
    wb.close()

    if first_col == 'DATA' and len(header_row) > 2:
        others = [str(h).strip().upper() for h in header_row[1:] if h is not None]
        if len(others) >= 1 and not any(h in ('VENDEDOR', 'VALOR', 'OBSERVACAO') for h in others[:3]):
            return 'matrix'

    return 'list'
