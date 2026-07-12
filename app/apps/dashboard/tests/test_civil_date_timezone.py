"""PROMPT 43 - datas civis (DateField) nao podem sofrer deslocamento de fuso.

Uma venda salva como 23/06/2026 deve aparecer como 23/06/2026 em backend,
JSON entregue ao template, PDF, CSV, XLSX e mobile. O bug real era o
JavaScript interpretar strings YYYY-MM-DD como UTC (new Date('2026-06-23'))
e recuar um dia no timezone America/Sao_Paulo.

Como o projeto nao possui runner JavaScript, cobrimos:
- o dado civil entregue ao template (JSON/contexto);
- a igualdade entre PDF/CSV/XLSX/mobile/detalhe;
- regressao no HTML/JS renderizado (uso do helper civil);
- o helper parseCivilDate construido por componentes locais;
- o DateField preservado sob TIME_ZONE=America/Sao_Paulo;
- o DateTimeField de auditoria mantendo conversao de timezone normal.
"""

import io
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from django.conf import settings
from django.test import TestCase, override_settings
from django.urls import reverse
from django.core.cache import cache
from rest_framework.test import APIClient

from app.apps.accounts.models import Tenant, User
from app.apps.sellers.models import Seller
from app.apps.commissions.models import CommissionPeriod, SellerCommission
from app.apps.sales.models import Sale


REPO_ROOT = Path(settings.BASE_DIR)
TEMPLATES_DIR = REPO_ROOT / 'templates'
STATIC_DIR = REPO_ROOT / 'static'


class CivilDateBackendTest(TestCase):
    """Datas civis entregues por API/PDF/CSV/XLSX sem deslocamento."""

    def setUp(self):
        cache.clear()
        self.tenant = Tenant.objects.create(
            company_name='Civil Date Co', slug='civil-date-co',
            default_commission_rate=Decimal('0.01'),
            period_start_day=21,
        )
        self.manager = User.objects.create_user(
            username='gestor_cd', password='pass123',
            role=User.Role.MANAGER, tenant=self.tenant,
        )
        self.seller_user = User.objects.create_user(
            username='vendedor_cd', password='pass123',
            role=User.Role.SELLER, tenant=self.tenant,
        )
        self.seller = Seller.objects.create(
            tenant=self.tenant, name='Vendedor Civil', phone='11999998888',
            user=self.seller_user, commission_rate=Decimal('0.01'),
        )
        # Competencia 21/06 a 20/07/2026 cobre 23/06, 30/06 e 01/07.
        self.period = CommissionPeriod.objects.create(
            tenant=self.tenant, month=7, year=2026,
            start_date=date(2026, 6, 21), end_date=date(2026, 7, 20),
            label='Julho/2026',
        )
        SellerCommission.objects.create(
            period=self.period, seller=self.seller,
            commission_rate=Decimal('0.01'),
            status=SellerCommission.Status.ABERTA,
        )
        # Datas de fronteira, historicamente as que "recuavam" um dia.
        self.sale_23 = Sale.objects.create(
            tenant=self.tenant, seller=self.seller, origin=Sale.Origin.MANUAL,
            amount=100000, sale_date='2026-06-23', created_by=self.seller_user,
        )
        self.sale_30 = Sale.objects.create(
            tenant=self.tenant, seller=self.seller, origin=Sale.Origin.MANUAL,
            amount=943672, sale_date='2026-06-30', created_by=self.seller_user,
        )
        self.sale_01 = Sale.objects.create(
            tenant=self.tenant, seller=self.seller, origin=Sale.Origin.MANUAL,
            amount=200000, sale_date='2026-07-01', created_by=self.seller_user,
        )

    def _auth(self):
        client = APIClient()
        resp = client.post(reverse('api-login'), {
            'username': self.manager.username, 'password': 'pass123',
        }, format='json')
        client.credentials(HTTP_AUTHORIZATION=f'Bearer {resp.data["access"]}')
        return client

    def _detail(self, client):
        # O nome 'api-seller-detail' colide com o router do SellerViewSet;
        # a view do gestor responde no caminho literal manager/seller/.
        return client.get(
            f'/api/manager/seller/{self.seller.uuid}/'
            + f'?period={self.period.uuid}',
        )

    def test_detail_delivers_civil_dates_unshifted(self):
        # LOTE 6 #1/#2/#3: 23/06, 30/06 e 01/07 permanecem intactas no JSON.
        client = self._auth()
        resp = self._detail(client)
        self.assertEqual(resp.status_code, 200)
        dates = {s['sale_date'] for s in resp.data['manual_sales']}
        self.assertIn('2026-06-23', dates)
        self.assertIn('2026-06-30', dates)
        self.assertIn('2026-07-01', dates)

    def test_detail_evolution_dates_unshifted(self):
        # O grafico consome evolution[].date (DateField civil).
        client = self._auth()
        resp = self._detail(client)
        evo_dates = {e['date'] for e in resp.data['evolution']}
        self.assertIn('2026-06-23', evo_dates)
        self.assertIn('2026-06-30', evo_dates)
        self.assertIn('2026-07-01', evo_dates)

    def test_pdf_and_detail_show_same_dates(self):
        # LOTE 6 #4: PDF exibe as mesmas datas do detalhe.
        client = self._auth()
        with patch('weasyprint.HTML') as html_class:
            html_class.return_value.write_pdf.return_value = b'%PDF-1.4'
            resp = client.get(
                reverse('api-seller-report-pdf', args=[self.seller.uuid])
                + f'?period={self.period.uuid}',
            )
        self.assertEqual(resp.status_code, 200)
        html = html_class.call_args.kwargs['string']
        self.assertIn('23/06/2026', html)
        self.assertIn('30/06/2026', html)
        self.assertIn('01/07/2026', html)

    def test_csv_shows_same_dates(self):
        # O CSV do relatorio usa ISO (YYYY-MM-DD): data civil, sem shift.
        client = self._auth()
        resp = client.get(
            reverse('api-seller-report-csv', args=[self.seller.uuid])
            + f'?period={self.period.uuid}',
        )
        content = resp.content.decode('utf-8')
        self.assertIn('2026-06-23', content)
        self.assertIn('2026-06-30', content)
        self.assertIn('2026-07-01', content)

    def test_xlsx_shows_same_dates(self):
        from openpyxl import load_workbook
        client = self._auth()
        resp = client.get(
            reverse('api-seller-report-xlsx', args=[self.seller.uuid])
            + f'?period={self.period.uuid}',
        )
        workbook = load_workbook(io.BytesIO(resp.content), data_only=True)
        rendered = '\n'.join(
            str(value)
            for row in workbook.active.iter_rows(values_only=True)
            for value in row if value
        )
        self.assertIn('23/06/2026', rendered)
        self.assertIn('30/06/2026', rendered)
        self.assertIn('01/07/2026', rendered)


class CivilDateMobileTest(TestCase):
    """Mobile entrega date/date_iso civis e filtro YYYY-MM correto."""

    def setUp(self):
        cache.clear()
        self.tenant = Tenant.objects.create(
            company_name='Mobile CD', slug='mobile-cd',
            default_commission_rate=Decimal('0.01'),
            period_start_day=21,
        )
        self.seller_user = User.objects.create_user(
            username='vendmob_cd', password='pass123',
            role=User.Role.SELLER, tenant=self.tenant,
        )
        self.seller = Seller.objects.create(
            tenant=self.tenant, name='Vend Mobile CD', phone='55999990000',
            user=self.seller_user, commission_rate=Decimal('0.01'),
        )
        self.period = CommissionPeriod.objects.create(
            tenant=self.tenant, month=7, year=2026,
            start_date=date(2026, 6, 21), end_date=date(2026, 7, 20),
            label='Julho/2026',
        )
        Sale.objects.create(
            tenant=self.tenant, seller=self.seller, origin=Sale.Origin.MANUAL,
            amount=100000, sale_date='2026-06-23', created_by=self.seller_user,
        )
        Sale.objects.create(
            tenant=self.tenant, seller=self.seller, origin=Sale.Origin.MANUAL,
            amount=200000, sale_date='2026-07-01', created_by=self.seller_user,
        )
        self.client.force_login(self.seller_user)

    def test_mobile_matches_detail_dates(self):
        # LOTE 6 #5: mobile entrega a mesma data do detalhe (23/06 e 01/07).
        resp = self.client.get('/dashboard/mobile/vendas/')
        self.assertEqual(resp.status_code, 200)
        by_iso = {s['date_iso']: s for s in resp.context['sales_json']}
        self.assertEqual(by_iso['2026-06-23']['date'], '23/06/2026')
        self.assertEqual(by_iso['2026-07-01']['date'], '01/07/2026')

    def test_mobile_month_filter_keys_correct(self):
        # LOTE 6 #8: filtro YYYY-MM continua correto (sem recuo de mes).
        resp = self.client.get('/dashboard/mobile/vendas/')
        keys = {o['key'] for o in resp.context['month_options']}
        self.assertIn('2026-06', keys)
        self.assertIn('2026-07', keys)
        month_keys = {s['month_key'] for s in resp.context['sales_json']}
        self.assertEqual(month_keys, {'2026-06', '2026-07'})


class CivilDateImportTest(TestCase):
    """Importacao: cada linha conserva a propria data; linha vazia nao cria."""

    def setUp(self):
        cache.clear()
        self.tenant = Tenant.objects.create(
            company_name='Import CD', slug='import-cd',
            default_commission_rate=Decimal('0.01'),
            period_start_day=21,
        )
        self.manager = User.objects.create_user(
            username='mgr_imp_cd', password='pass123',
            role=User.Role.MANAGER, tenant=self.tenant,
        )
        self.seller_user = User.objects.create_user(
            username='sel_imp_cd', password='pass123',
            role=User.Role.SELLER, tenant=self.tenant,
        )
        self.seller = Seller.objects.create(
            tenant=self.tenant, name='Carlos Silva', phone='11911111111',
            user=self.seller_user, commission_rate=Decimal('0.01'),
        )
        self.period = CommissionPeriod.objects.create(
            tenant=self.tenant, month=7, year=2026,
            start_date=date(2026, 6, 21), end_date=date(2026, 7, 20),
            status=CommissionPeriod.Status.ABERTA,
        )
        SellerCommission.objects.create(
            period=self.period, seller=self.seller,
            commission_rate=Decimal('0.01'),
            status=SellerCommission.Status.ABERTA,
        )

    def _auth(self):
        client = APIClient()
        resp = client.post(reverse('api-login'), {
            'username': self.manager.username, 'password': 'pass123',
        }, format='json')
        client.credentials(HTTP_AUTHORIZATION=f'Bearer {resp.data["access"]}')
        return client

    def _preview(self, client, content, filename='vendas.csv'):
        file_obj = io.BytesIO(content.encode('utf-8'))
        file_obj.name = filename
        return client.post(
            reverse('api-sale-import-preview'),
            {'file': file_obj}, format='multipart',
        )

    def test_preview_preserves_each_row_date(self):
        # LOTE 6 #9: 29/06 (FALTA, sem valor) e 30/06 (com valor).
        csv = '\n'.join([
            'data;vendedor;valor',
            '29/06/2026;Carlos Silva;FALTA',
            '30/06/2026;Carlos Silva;9.436,72',
        ])
        client = self._auth()
        resp = self._preview(client, csv)
        self.assertEqual(resp.status_code, 200)
        results = resp.data['results']
        # A linha 30/06 conserva sua propria data, nunca 29/06.
        ok_rows = [r for r in results if r['status'] == 'ok']
        self.assertEqual(len(ok_rows), 1)
        self.assertEqual(ok_rows[0]['sale_date'], '2026-06-30')
        self.assertEqual(ok_rows[0]['amount_cents'], 943672)
        # A linha 29/06 sem valor valido nao pode virar venda.
        error_rows = [r for r in results if r['status'] == 'error']
        self.assertTrue(any(r['sale_date'] == '2026-06-29' for r in error_rows))

    def test_confirm_empty_line_does_not_create_sale(self):
        # LOTE 6 #10/#11: valor do dia 30 nunca associado ao dia 29.
        csv = '\n'.join([
            'data;vendedor;valor',
            '29/06/2026;Carlos Silva;FALTA',
            '30/06/2026;Carlos Silva;9.436,72',
        ])
        client = self._auth()
        preview = self._preview(client, csv)
        rows = [
            {
                'row_number': r['row_number'],
                'resolved_seller_uuid': r['resolved_seller_uuid'],
                'amount_cents': r['amount_cents'],
                'sale_date': r['sale_date'],
                'notes': r['notes'],
            }
            for r in preview.data['results'] if r['status'] == 'ok'
        ]
        resp = client.post(reverse('api-sale-import-confirm'), {
            'rows': rows, 'filename': 'vendas.csv',
            'file_hash': preview.data['file_hash'],
            'force_reimport': False,
        }, format='json')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data['created'], 1)
        self.assertFalse(
            Sale.objects.filter(sale_date=date(2026, 6, 29)).exists(),
        )
        sale = Sale.objects.get(sale_date=date(2026, 6, 30))
        self.assertEqual(sale.amount, 943672)


@override_settings(TIME_ZONE='America/Sao_Paulo', USE_TZ=True)
class CivilDateTimezoneModelTest(TestCase):
    """DateField civil nao muda sob America/Sao_Paulo; DateTime muda normal."""

    def setUp(self):
        cache.clear()
        self.tenant = Tenant.objects.create(
            company_name='TZ CD', slug='tz-cd',
            default_commission_rate=Decimal('0.01'),
        )
        self.seller_user = User.objects.create_user(
            username='tz_seller', password='pass123',
            role=User.Role.SELLER, tenant=self.tenant,
        )
        self.seller = Seller.objects.create(
            tenant=self.tenant, name='TZ Seller', phone='11900000000',
            user=self.seller_user, commission_rate=Decimal('0.01'),
        )

    def test_datefield_not_shifted_by_timezone(self):
        # LOTE 6 #12.
        sale = Sale.objects.create(
            tenant=self.tenant, seller=self.seller, origin=Sale.Origin.MANUAL,
            amount=100000, sale_date='2026-06-23', created_by=self.seller_user,
        )
        sale.refresh_from_db()
        self.assertEqual(sale.sale_date, date(2026, 6, 23))
        self.assertEqual(sale.sale_date.isoformat(), '2026-06-23')
        # weekday civil de 23/06/2026 = terca-feira (Python weekday()==1).
        self.assertEqual(sale.sale_date.weekday(), 1)

    def test_datetimefield_audit_uses_timezone(self):
        # LOTE 6 #13: created_at (DateTimeField) permanece aware/tz-normal.
        sale = Sale.objects.create(
            tenant=self.tenant, seller=self.seller, origin=Sale.Origin.MANUAL,
            amount=100000, sale_date='2026-06-23', created_by=self.seller_user,
        )
        self.assertIsNotNone(sale.created_at.tzinfo)


class CivilDateTemplateRegressionTest(TestCase):
    """Regressao no HTML/JS: uso do helper civil, sem new Date(iso) cru."""

    def _read(self, *parts):
        return (TEMPLATES_DIR.joinpath(*parts)).read_text(encoding='utf-8')

    def test_helper_parses_by_local_components(self):
        # LOTE 6 #6: o helper constroi a data por componentes locais, o que
        # garante o dia da semana civil correto (sem interpretacao UTC).
        src = (STATIC_DIR / 'js' / 'format.js').read_text(encoding='utf-8')
        self.assertIn('function parseCivilDate', src)
        self.assertIn('new Date(year, month - 1, day)', src)
        self.assertIn('function formatCivilDate', src)

    def test_vendedor_detalhe_uses_civil_helper(self):
        html = self._read('dashboard', 'gestor', 'vendedor_detalhe.html')
        self.assertNotIn("new Date(s.sale_date)", html)
        self.assertNotIn("new Date(e.date)", html)
        self.assertIn('formatCivilDate(s.sale_date)', html)
        self.assertIn('formatCivilDate(e.date', html)
        # DateTimeField de auditoria mantem conversao normal de timezone.
        self.assertIn('new Date(entry.changed_at)', html)

    def test_importar_vendas_uses_civil_helper(self):
        html = self._read('dashboard', 'gestor', 'importar_vendas.html')
        self.assertNotIn("new Date(r.sale_date", html)
        self.assertIn('formatCivilDate(r.sale_date)', html)

    def test_minhas_vendas_uses_civil_helper(self):
        # LOTE 6 #7: agrupamento semanal usa a data civil (sem deslocamento).
        html = self._read('mobile', 'minhas_vendas.html')
        self.assertNotIn("new Date(s.date_iso", html)
        self.assertIn('parseCivilDate(s.date_iso)', html)
